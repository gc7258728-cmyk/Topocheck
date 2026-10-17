# -*- coding: utf-8 -*-
"""
TopoCheck — Application pédagogique de nivellement géométrique direct,
calcul d'altitudes, contrôle de fermeture, compensation et analyse d'incertitudes.

Public cible : Étudiants en 3e année de Génie Civil (ENSPY — Yaoundé).
Question centrale : « Mes lectures de nivellement permettent-elles de retrouver la bonne altitude ? »
Fonctionnement 100 % local, sans connexion externe, sans clé API ni service payant.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import copy
from html import escape
from pathlib import Path
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import plotly.graph_objects as go
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    Image,
    PageBreak,
    KeepTogether,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from scipy import stats

try:
    import streamlit as st
except ImportError:  # Permet l'import direct des fonctions pour les tests unitaires hors Streamlit
    st = None


# ==============================================================================
# 1. PALETTE & IDENTITÉ VISUELLE « CARNET DE TERRAIN CONTEMPORAIN »
# ==============================================================================
# Direction visuelle :
# - Fond gris perle (#F3F4F6), cartes blanches (#FFFFFF), textes anthracite (#1E2229)
# - Couleur principale : Framboise (#B81D5A)
# - Accents : Jaune signalétique (#E59A00) et Turquoise (#0D9488)
# - Aucun fond beige, aucun dégradé décoratif, aucun bandeau vert/bleu ni 4 cartes identiques.

PALETTE = {
    "bg_pearl": "#F3F4F6",
    "bg_sidebar": "#E8EAEE",
    "surface": "#FFFFFF",
    "border": "#D5D9E0",
    "anthracite": "#1E2229",
    "text_muted": "#4A515D",
    "raspberry": "#B81D5A",
    "raspberry_soft": "#FDF2F6",
    "turquoise": "#0D9488",
    "turquoise_soft": "#F0FDFA",
    "signal_yellow": "#E59A00",
    "signal_yellow_soft": "#FFFBEB",
    "danger": "#B91C1C",
    "success": "#0F766E",
}

RUBRIQUES = [
    "📐 Tableau de bord",
    "📓 Carnet de nivellement",
    "⚖️ Fermeture et compensation",
    "🎲 Incertitudes",
    "🔍 Comparaison",
    "💾 Données et rapport",
    "🧪 Méthode et tests",
]


# ==============================================================================
# 2. JEUX DE DONNÉES DE RÉFÉRENCE & EXEMPLES PÉDAGOGIQUES
# ==============================================================================

def get_preset_datasets() -> Dict[str, Dict[str, Any]]:
    """Retourne les scénarios pédagogiques et les cas de tests de référence."""
    return {
        "chantier_enspy": {
            "label": "Campus ENSPY — Cheminement encadré (4 stations, avec longueurs)",
            "description": (
                "Cheminement pédagogique entre le repère RN_Entree (100,000 m) et le repère "
                "RN_LaboGC (101,150 m) avec 4 stations, incluant une descente intermédiaire "
                "et une fermeture de +12,0 mm."
            ),
            "project_name": "Levé Campus ENSPY — Bloc Génie Civil",
            "operator": "Groupe TP 3GC — ENSPY",
            "z0": 100.000,
            "regime": "encadre",
            "z_ref_final": 101.150,
            "tolerance_mm": 15.0,
            "comp_mode": "longueurs",
            "sigma_mm": 2.0,
            "sigma_alt_mm": 0.8,
            "n_draws": 4000,
            "seed": 42,
            "stations": [
                {
                    "station_id": "S1",
                    "point_ar": "RN_Entree",
                    "point_av": "P1",
                    "lecture_ar_m": 1.842,
                    "lecture_av_m": 1.120,
                    "longueur_m": 45.0,
                },
                {
                    "station_id": "S2",
                    "point_ar": "P1",
                    "point_av": "P2",
                    "lecture_ar_m": 1.615,
                    "lecture_av_m": 0.985,
                    "longueur_m": 60.0,
                },
                {
                    "station_id": "S3",
                    "point_ar": "P2",
                    "point_av": "P3",
                    "lecture_ar_m": 1.110,
                    "lecture_av_m": 1.540,
                    "longueur_m": 50.0,
                },
                {
                    "station_id": "S4",
                    "point_ar": "P3",
                    "point_av": "RN_LaboGC",
                    "lecture_ar_m": 1.790,
                    "lecture_av_m": 1.550,
                    "longueur_m": 45.0,
                },
            ],
        },
        "test_simple_1_station": {
            "label": "Test de référence 1 — 1 station (Départ 100 m, AR 1,50 m, AV 1,20 m)",
            "description": (
                "Cas de référence du cahier des charges : Z0 = 100,000 m, AR = 1,500 m, "
                "AV = 1,200 m -> Dénivelée +0,300 m, Z_fin = 100,300 m. "
                "Avec référence 100,280 m : fermeture +0,020 m (+20 mm), correction -20 mm."
            ),
            "project_name": "Test de référence 1 — Station unique",
            "operator": "Validation TopoCheck",
            "z0": 100.000,
            "regime": "encadre",
            "z_ref_final": 100.280,
            "tolerance_mm": 25.0,
            "comp_mode": "egaux",
            "sigma_mm": 2.0,
            "sigma_alt_mm": 1.0,
            "n_draws": 3000,
            "seed": 42,
            "stations": [
                {
                    "station_id": "S1",
                    "point_ar": "RN0",
                    "point_av": "RN1",
                    "lecture_ar_m": 1.500,
                    "lecture_av_m": 1.200,
                    "longueur_m": 50.0,
                }
            ],
        },
        "test_2_stations_boucle": {
            "label": "Test de référence 2 & 3 — 2 stations (Arrivée 100 m, sigma=2 mm -> 4 mm)",
            "description": (
                "Cas de référence à 2 stations : S1 (AR 1,50 m, AV 1,20 m) et S2 (AR 1,10 m, AV 1,40 m). "
                "Arrivée à 100,000 m, fermeture nulle. Avec 4 lectures indépendantes de 2 mm, "
                "l'écart-type théorique final vaut exactement 4,0 mm."
            ),
            "project_name": "Test de référence 2 & 3 — Deux stations",
            "operator": "Validation TopoCheck",
            "z0": 100.000,
            "regime": "boucle",
            "z_ref_final": 100.000,
            "tolerance_mm": 8.0,
            "comp_mode": "egaux",
            "sigma_mm": 2.0,
            "sigma_alt_mm": 1.0,
            "n_draws": 5000,
            "seed": 12345,
            "stations": [
                {
                    "station_id": "S1",
                    "point_ar": "RN_A",
                    "point_av": "P1",
                    "lecture_ar_m": 1.500,
                    "lecture_av_m": 1.200,
                    "longueur_m": 40.0,
                },
                {
                    "station_id": "S2",
                    "point_ar": "P1",
                    "point_av": "RN_A",
                    "lecture_ar_m": 1.100,
                    "lecture_av_m": 1.400,
                    "longueur_m": 60.0,
                },
            ],
        },
        "cheminement_ouvert_descente": {
            "label": "Exemple — Cheminement ouvert en descente (sans référence finale)",
            "description": (
                "Antenne topographique en descente vers un exutoire pluvial sans repère final connu : "
                "illustre que les dénivelées négatives sont normales et que la fermeture n'est pas calculable."
            ),
            "project_name": "Antenne Pluviale — Cheminement Ouvert",
            "operator": "Groupe TP 3GC — ENSPY",
            "z0": 725.400,
            "regime": "ouvert",
            "z_ref_final": None,
            "tolerance_mm": 12.0,
            "comp_mode": "egaux",
            "sigma_mm": 1.5,
            "sigma_alt_mm": 0.7,
            "n_draws": 3000,
            "seed": 99,
            "stations": [
                {
                    "station_id": "S1",
                    "point_ar": "RN_Haut",
                    "point_av": "Regard_1",
                    "lecture_ar_m": 0.850,
                    "lecture_av_m": 1.920,
                    "longueur_m": 35.0,
                },
                {
                    "station_id": "S2",
                    "point_ar": "Regard_1",
                    "point_av": "Regard_2",
                    "lecture_ar_m": 0.910,
                    "lecture_av_m": 2.140,
                    "longueur_m": 45.0,
                },
                {
                    "station_id": "S3",
                    "point_ar": "Regard_2",
                    "point_av": "Exutoire",
                    "lecture_ar_m": 1.050,
                    "lecture_av_m": 1.880,
                    "longueur_m": 40.0,
                },
            ],
        },
    }


# ==============================================================================
# 3. VALIDATION ET CALCULS TOPOGRAPHIQUES (CARNET, FERMETURE, COMPENSATION)
# ==============================================================================

REQUIRED_COLUMNS = [
    "station_id",
    "point_ar",
    "point_av",
    "lecture_ar_m",
    "lecture_av_m",
]
OPTIONAL_COLUMNS = ["longueur_m"]
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


def finite_number(value: Any, low: float, high: float) -> bool:
    return (not isinstance(value, (bool, np.bool_)) and isinstance(value, (int, float, np.number))
            and math.isfinite(float(value)) and low <= value <= high)


def exceeds_tolerance(values_mm: Any, tolerance_mm: float) -> Any:
    # Stabilise only sub-nanometre numerical noise at a strict threshold.
    return np.round(np.abs(values_mm), 6) > round(float(tolerance_mm), 6)


def normalize_stations_df(stations_input: Any) -> pd.DataFrame:
    """Convertit une liste de dictionnaires ou un DataFrame en DataFrame normalisé."""
    issues = list(stations_input.attrs.get('input_errors', [])) if isinstance(stations_input, pd.DataFrame) else []
    try:
        df = stations_input.copy() if isinstance(stations_input, pd.DataFrame) else pd.DataFrame(stations_input)
    except (TypeError, ValueError):
        df = pd.DataFrame(columns=ALL_COLUMNS)
        issues.append('Le carnet doit être un tableau de stations.')

    for col in REQUIRED_COLUMNS:
        if col not in df.columns:
            issues.append(f'Colonne obligatoire absente : {col}.')
            df[col] = "" if col in ("station_id", "point_ar", "point_av") else np.nan
    if "longueur_m" not in df.columns:
        df["longueur_m"] = np.nan

    df = df[ALL_COLUMNS].copy()
    for str_col in ("station_id", "point_ar", "point_av"):
        df[str_col] = df[str_col].fillna("").astype(str).str.strip()

    for num_col in ("lecture_ar_m", "lecture_av_m", "longueur_m"):
        raw = df[num_col]
        missing = raw.isna() | raw.astype(str).str.strip().eq('')
        converted = pd.to_numeric(raw, errors="coerce")
        if ((~missing) & ~np.isfinite(converted)).any() or raw.map(lambda v:isinstance(v,(bool,np.bool_))).any():
            issues.append(f'{num_col} : valeur non numérique, booléenne ou infinie.')
        df[num_col] = converted

    df = df.reset_index(drop=True)
    df.attrs['input_errors'] = list(dict.fromkeys(issues))
    return df


def compute_state_signature(
    df: pd.DataFrame,
    z0: float,
    regime: str,
    z_ref_final: Optional[float],
    tolerance_mm: float,
    comp_mode: str,
    sigma_mm: float,
    n_draws: int,
    seed: int,
) -> str:
    """Calcule une empreinte SHA-256 déterministe de l'état des données et paramètres."""
    payload = {
        "z0": float(z0) if z0 is not None and not pd.isna(z0) else None,
        "regime": str(regime),
        "z_ref_final": (
            float(z_ref_final)
            if z_ref_final is not None and not pd.isna(z_ref_final)
            else None
        ),
        "tolerance_mm": float(tolerance_mm),
        "comp_mode": str(comp_mode),
        "sigma_mm": float(sigma_mm),
        "n_draws": int(n_draws),
        "seed": int(seed),
        "rows": df.astype(object).where(pd.notna(df), None).to_dict(orient="records"),
    }
    raw = json.dumps(payload, sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]


def validate_carnet(
    stations_input: Any,
    z0: float,
    regime: str = "encadre",
    z_ref_final: Optional[float] = None,
) -> Dict[str, Any]:
    """
    Valide le carnet de nivellement :
    - présence des colonnes et d'au moins une station ;
    - identifiants uniques et non vides ;
    - continuité de la chaîne de points (point_ar[i] == point_av[i-1]) ;
    - validité des lectures sur mire (valeurs numériques positives) ;
    - cohérence des longueurs de station (optionnelles mais positives si renseignées) ;
    - cohérence du régime de fermeture (ouvert, encadré, boucle).
    """
    errors: List[str] = []
    warnings: List[str] = []

    df = normalize_stations_df(stations_input)
    errors.extend(df.attrs.get('input_errors', []))
    if len(df)>100:errors.append('Le carnet est limité à 100 stations.')
    if regime not in ('encadre','boucle','ouvert'):errors.append('Type de cheminement inconnu.')

    if not finite_number(z0, -500, 10000):
        errors.append("L'altitude de départ Z0 est invalide ou manquante.")

    if len(df) == 0:
        errors.append("Le carnet ne contient aucune station. Ajoutez au moins une station.")
        return {
            "is_valid": False,
            "errors": errors,
            "warnings": warnings,
            "all_lengths_valid": False,
            "continuity_ok": False,
            "loop_point_matches": False,
            "df": df,
        }

    seen_ids = set()
    for idx, row in df.iterrows():
        row_num = idx + 1
        sid = row["station_id"]
        p_ar = row["point_ar"]
        p_av = row["point_av"]
        if any(len(v)>40 for v in (sid,p_ar,p_av)):
            errors.append(f'Ligne {row_num} : identifiants limités à 40 caractères.')

        if not sid:
            errors.append(f"Ligne {row_num} : l'identifiant de station est vide.")
        elif sid in seen_ids:
            errors.append(f"Ligne {row_num} : l'identifiant de station '{sid}' est dupliqué.")
        else:
            seen_ids.add(sid)

        if not p_ar:
            errors.append(f"Station {sid or row_num} : le nom du point arrière est vide.")
        if not p_av:
            errors.append(f"Station {sid or row_num} : le nom du point avant est vide.")
        if p_ar and p_av and p_ar == p_av:
            errors.append(
                f"Station {sid or row_num} : le point arrière et le point avant sont identiques ('{p_ar}')."
            )

        l_ar = row["lecture_ar_m"]
        l_av = row["lecture_av_m"]
        if pd.isna(l_ar) or not np.isfinite(l_ar):
            errors.append(f"Station {sid or row_num} : la lecture arrière est manquante ou non numérique.")
        elif l_ar < 0 or l_ar > 10:
            errors.append(
                f"Station {sid or row_num} : lecture arrière hors plage de saisie 0 à 10 m. La dénivelée peut être négative."
            )
        elif l_ar == 0 or l_ar > 5.0:
            warnings.append(
                f"Station {sid or row_num} : lecture arrière inhabituelle ({l_ar:.3f} m) pour une mire standard."
            )

        if pd.isna(l_av) or not np.isfinite(l_av):
            errors.append(f"Station {sid or row_num} : la lecture avant est manquante ou non numérique.")
        elif l_av < 0 or l_av > 10:
            errors.append(
                f"Station {sid or row_num} : lecture avant hors plage de saisie 0 à 10 m."
            )
        elif l_av == 0 or l_av > 5.0:
            warnings.append(
                f"Station {sid or row_num} : lecture avant inhabituelle ({l_av:.3f} m) pour une mire standard."
            )

        long_m = row["longueur_m"]
        if not pd.isna(long_m):
            if not np.isfinite(long_m) or long_m <= 0 or long_m > 10000:
                errors.append(
                    f"Station {sid or row_num} : longueur requise strictement positive et au plus 10 000 m, ou cellule vide."
                )
            elif long_m > 250.0:
                warnings.append(
                    f"Station {sid or row_num} : longueur de cheminement élevée ({long_m:.1f} m). "
                    "Vérifier la saisie ; cette longueur n'est pas une portée de visée."
                )

    continuity_ok = True
    for idx in range(1, len(df)):
        prev_av = df.loc[idx - 1, "point_av"]
        curr_ar = df.loc[idx, "point_ar"]
        curr_sid = df.loc[idx, "station_id"] or f"#{idx + 1}"
        prev_sid = df.loc[idx - 1, "station_id"] or f"#{idx}"
        if prev_av != curr_ar:
            continuity_ok = False
            errors.append(
                f"Rupture de continuité entre {prev_sid} et {curr_sid} : "
                f"le point avant de {prev_sid} est '{prev_av}', mais le point arrière de {curr_sid} est '{curr_ar}'."
            )

    lengths_series = df["longueur_m"]
    all_lengths_valid = bool(
        lengths_series.notna().all() and (lengths_series > 0).all()
    )
    if lengths_series.notna().any() and not all_lengths_valid:
        warnings.append(
            "Certaines longueurs de stations sont vides ou nulles : la compensation proportionnelle "
            "aux longueurs nécessite une longueur strictement positive sur chaque station."
        )

    first_point = df.loc[0, "point_ar"]
    last_point = df.loc[len(df) - 1, "point_av"]
    loop_point_matches = bool(first_point and last_point and first_point == last_point)

    if regime == "encadre":
        if not finite_number(z_ref_final,-500,10000):
            errors.append(
                "En cheminement encadré entre deux repères connus, l'altitude finale de référence doit être renseignée."
            )
    elif regime == "boucle":
        if not loop_point_matches:
            errors.append(
                f"Mode 'Boucle fermée' sélectionné, mais le point d'arrivée ('{last_point}') "
                f"ne porte pas le même nom que le point de départ ('{first_point}'). "
                "Corriger le point final ou choisir un autre type de cheminement."
            )

    return {
        "is_valid": len(errors) == 0,
        "errors": errors,
        "warnings": warnings,
        "all_lengths_valid": all_lengths_valid,
        "continuity_ok": continuity_ok,
        "loop_point_matches": loop_point_matches,
        "df": df,
    }


def compute_leveling(
    stations_input: Any,
    z0: float,
    regime: str = "encadre",
    z_ref_final: Optional[float] = None,
    tolerance_mm: float = 15.0,
    comp_mode: str = "egaux",
) -> Dict[str, Any]:
    """
    Exécute l'intégralité des calculs de nivellement géométrique.
    """
    validation = validate_carnet(stations_input, z0, regime, z_ref_final)
    df = validation["df"].copy()
    if not finite_number(tolerance_mm,.1,1000):validation['errors'].append('Tolérance : nombre entre 0,1 et 1 000 mm.')
    if comp_mode not in ('egaux','longueurs'):validation['errors'].append('Méthode de compensation inconnue.')
    if comp_mode=='longueurs' and not validation['all_lengths_valid']:
        validation['errors'].append('Compensation par longueurs : renseigner une longueur positive par station ; aucune bascule automatique.')
    validation['is_valid']=not validation['errors']

    if not validation["is_valid"]:
        return {
            "validation": validation,
            "is_valid": False,
            "stations_table": df,
            "points_table": pd.DataFrame(),
        }

    z0_val = float(z0)
    n_stations = len(df)

    df["denivelee_m"] = df["lecture_ar_m"] - df["lecture_av_m"]
    df["denivelee_mm"] = df["denivelee_m"] * 1000.0
    df["sens"] = np.where(
        df["denivelee_m"] > 1e-9,
        "Montée",
        np.where(df["denivelee_m"] < -1e-9, "Descente", "Palier"),
    )

    z_ar_list: List[float] = []
    hv_list: List[float] = []
    z_av_brut_list: List[float] = []

    curr_z = z0_val
    for _, row in df.iterrows():
        z_ar_list.append(curr_z)
        hv = curr_z + float(row["lecture_ar_m"])
        hv_list.append(hv)
        curr_z = curr_z + float(row["denivelee_m"])
        z_av_brut_list.append(curr_z)

    df["z_ar_brut_m"] = z_ar_list
    df["plan_visee_m"] = hv_list
    df["z_av_brut_m"] = z_av_brut_list

    sum_ar_m = float(df["lecture_ar_m"].sum())
    sum_av_m = float(df["lecture_av_m"].sum())
    diff_lectures_m = sum_ar_m - sum_av_m
    sum_denivelees_m = float(df["denivelee_m"].sum())
    z_fin_brut_m = float(z_av_brut_list[-1])
    diff_altitudes_m = z_fin_brut_m - z0_val

    arith_error_m = max(
        abs(diff_lectures_m - sum_denivelees_m),
        abs(sum_denivelees_m - diff_altitudes_m),
    )
    arithmetic_check_ok = arith_error_m < 1e-9

    if regime == "ouvert":
        z_ref_used: Optional[float] = None
        fermeture_m: Optional[float] = None
        fermeture_mm: Optional[float] = None
        is_within_tolerance: Optional[bool] = None
        can_compensate = False
    elif regime == "boucle":
        z_ref_used = z0_val
        fermeture_m = z_fin_brut_m - z_ref_used
        fermeture_mm = fermeture_m * 1000.0
        is_within_tolerance = not bool(exceeds_tolerance(fermeture_mm,tolerance_mm))
        can_compensate = True
    else:  # "encadre"
        z_ref_used = float(z_ref_final)
        fermeture_m = z_fin_brut_m - z_ref_used
        fermeture_mm = fermeture_m * 1000.0
        is_within_tolerance = not bool(exceeds_tolerance(fermeture_mm,tolerance_mm))
        can_compensate = True

    effective_comp_mode = comp_mode
    comp_warning: Optional[str] = None
    if can_compensate and fermeture_m is not None:
        if comp_mode == "longueurs":
            if validation["all_lengths_valid"]:
                weights = df["longueur_m"].to_numpy(dtype=float)
            else:
                effective_comp_mode = "egaux"
                comp_warning = (
                    "Longueurs incomplètes ou nulles : bascule automatique sur des poids égaux "
                    "(1 par station) pour la compensation."
                )
                weights = np.ones(n_stations, dtype=float)
        else:
            weights = np.ones(n_stations, dtype=float)

        sum_weights = float(np.sum(weights))
        weight_ratios = weights / sum_weights
        corrections_m = -fermeture_m * weight_ratios
        corrections_mm = corrections_m * 1000.0
        denivelees_comp_m = df["denivelee_m"].to_numpy(dtype=float) + corrections_m

        z_ar_comp_list: List[float] = []
        z_av_comp_list: List[float] = []
        curr_z_comp = z0_val
        for dz_c in denivelees_comp_m:
            z_ar_comp_list.append(curr_z_comp)
            curr_z_comp = curr_z_comp + float(dz_c)
            z_av_comp_list.append(curr_z_comp)

        df["poids"] = weights
        df["part_poids_pct"] = weight_ratios * 100.0
        df["correction_m"] = corrections_m
        df["correction_mm"] = corrections_mm
        df["denivelee_comp_m"] = denivelees_comp_m
        df["z_ar_comp_m"] = z_ar_comp_list
        df["z_av_comp_m"] = z_av_comp_list

        sum_corrections_m = float(np.sum(corrections_m))
        sum_corrections_mm = float(np.sum(corrections_mm))
        z_fin_comp_m = float(z_av_comp_list[-1])
        comp_closure_error_m = abs((z_fin_comp_m - z_ref_used))
        comp_sum_check_ok = abs(sum_corrections_m + fermeture_m) < 1e-9 and comp_closure_error_m < 1e-9
    else:
        df["poids"] = np.nan
        df["part_poids_pct"] = np.nan
        df["correction_m"] = np.nan
        df["correction_mm"] = np.nan
        df["denivelee_comp_m"] = np.nan
        df["z_ar_comp_m"] = np.nan
        df["z_av_comp_m"] = np.nan
        sum_corrections_m = None
        sum_corrections_mm = None
        z_fin_comp_m = None
        comp_sum_check_ok = None

    points_rows: List[Dict[str, Any]] = [
        {
            "rang": 0,
            "point": df.loc[0, "point_ar"],
            "station_origine": "Départ",
            "distance_station_m": 0.0,
            "distance_cumulee_m": 0.0 if validation["all_lengths_valid"] else np.nan,
            "z_brut_m": z0_val,
            "z_comp_m": z0_val if can_compensate else np.nan,
            "ecart_comp_mm": 0.0 if can_compensate else np.nan,
        }
    ]
    cum_dist = 0.0
    for idx, row in df.iterrows():
        d_i = float(row["longueur_m"]) if validation["all_lengths_valid"] else np.nan
        if validation["all_lengths_valid"]:
            cum_dist += d_i
        z_b = float(row["z_av_brut_m"])
        z_c = float(row["z_av_comp_m"]) if can_compensate else np.nan
        points_rows.append(
            {
                "rang": idx + 1,
                "point": row["point_av"],
                "station_origine": row["station_id"],
                "distance_station_m": d_i,
                "distance_cumulee_m": cum_dist if validation["all_lengths_valid"] else np.nan,
                "z_brut_m": z_b,
                "z_comp_m": z_c,
                "ecart_comp_mm": (z_c - z_b) * 1000.0 if can_compensate else np.nan,
            }
        )
    points_df = pd.DataFrame(points_rows)

    total_length_m = float(df["longueur_m"].sum()) if validation["all_lengths_valid"] else None

    return {
        "validation": validation,
        "is_valid": True,
        "regime": regime,
        "z0_m": z0_val,
        "z_ref_used_m": z_ref_used,
        "n_stations": n_stations,
        "total_length_m": total_length_m,
        "all_lengths_valid": validation["all_lengths_valid"],
        "sum_ar_m": sum_ar_m,
        "sum_av_m": sum_av_m,
        "diff_lectures_m": diff_lectures_m,
        "sum_denivelees_m": sum_denivelees_m,
        "diff_altitudes_m": diff_altitudes_m,
        "z_fin_brut_m": z_fin_brut_m,
        "arithmetic_check_ok": arithmetic_check_ok,
        "arith_error_m": arith_error_m,
        "can_compensate": can_compensate,
        "fermeture_m": fermeture_m,
        "fermeture_mm": fermeture_mm,
        "tolerance_mm": float(tolerance_mm),
        "is_within_tolerance": is_within_tolerance,
        "comp_mode_requested": comp_mode,
        "comp_mode_effective": effective_comp_mode,
        "comp_warning": comp_warning,
        "sum_corrections_m": sum_corrections_m,
        "sum_corrections_mm": sum_corrections_mm,
        "z_fin_comp_m": z_fin_comp_m,
        "comp_sum_check_ok": comp_sum_check_ok,
        "stations_table": df,
        "points_table": points_df,
    }


# ==============================================================================
# 4. INCERTITUDES, PROPAGATION ANALYTIQUE & SIMULATION MONTE CARLO
# ==============================================================================

def binomial_ci_95(k: int, n: int) -> Tuple[float, float]:
    """
    Calcule l'intervalle de confiance binomial exact à 95 % (Clopper-Pearson)
    pour k succès sur n tirages indépendants via scipy.stats.beta.
    """
    if not finite_number(n,1,50000) or not finite_number(k,0,50000) or int(n)!=n or int(k)!=k or k>n:
        raise ValueError('Intervalle binomial : entiers 0 ≤ k ≤ n et n > 0 requis.')
    alpha = 0.05
    low = 0.0 if k <= 0 else float(stats.beta.ppf(alpha / 2.0, k, n - k + 1))
    high = 1.0 if k >= n else float(stats.beta.ppf(1.0 - alpha / 2.0, k + 1, n - k))
    return (low, high)


def wilson_ci_95_vectorized(k_arr: np.ndarray, n_arr: np.ndarray) -> Tuple[np.ndarray, np.ndarray]:
    """Intervalle binomial de Wilson à 95 % vectorisé pour la courbe de convergence."""
    z = 1.959963984540054
    p_hat = k_arr / n_arr
    denom = 1.0 + (z**2) / n_arr
    center = (p_hat + (z**2) / (2.0 * n_arr)) / denom
    half = (z * np.sqrt((p_hat * (1.0 - p_hat) + (z**2) / (4.0 * n_arr)) / n_arr)) / denom
    low = np.clip(center - half, 0.0, 1.0)
    high = np.clip(center + half, 0.0, 1.0)
    return low, high


def simulate_uncertainties(
    stations_input: Any,
    z0: float,
    sigma_mm: float = 2.0,
    n_draws: int = 3000,
    seed: int = 42,
    regime: str = "encadre",
    z_ref_final: Optional[float] = None,
    tolerance_mm: float = 15.0,
) -> Dict[str, Any]:
    """
    Simule des erreurs aléatoires normales centrées indépendantes sur chaque lecture (AR et AV).
    """
    if not finite_number(sigma_mm,0,50):raise ValueError('Écart-type : valeur entre 0 et 50 mm.')
    if not finite_number(n_draws,100,50000) or int(n_draws)!=n_draws:raise ValueError('Tirages : entier de 100 à 50 000.')
    if not finite_number(seed,0,2147483647) or int(seed)!=seed:raise ValueError('Graine : entier de 0 à 2 147 483 647.')
    base = compute_leveling(
        stations_input=stations_input,
        z0=z0,
        regime=regime,
        z_ref_final=z_ref_final,
        tolerance_mm=tolerance_mm,
    )
    if not base["is_valid"]:
        return {"is_valid": False, "base": base}

    df = base["stations_table"]
    n_stations = int(base["n_stations"])
    n_lectures = 2 * n_stations
    sigma_mm_val = float(sigma_mm)
    sigma_m_val = sigma_mm_val / 1000.0
    n_draws_val = int(n_draws)
    seed_val = int(seed)

    var_station_mm2 = 2.0 * (sigma_mm_val**2)
    sigma_station_mm = math.sqrt(var_station_mm2)
    var_total_mm2 = n_stations * var_station_mm2
    sigma_total_analytic_mm = math.sqrt(var_total_mm2)
    sigma_total_analytic_m = sigma_total_analytic_mm / 1000.0

    l_ar_nom = df["lecture_ar_m"].to_numpy(dtype=float)
    l_av_nom = df["lecture_av_m"].to_numpy(dtype=float)
    z0_val = float(base["z0_m"])
    z_fin_nom_m = float(base["z_fin_brut_m"])

    rng = np.random.default_rng(seed_val)
    # O(N) memory, with shared standard draws across precision scenarios.
    ecart_nominal_mm = np.zeros(n_draws_val)
    for _ in range(n_stations):
        ecart_nominal_mm += (rng.standard_normal(n_draws_val)-rng.standard_normal(n_draws_val))*sigma_mm_val
    z_fin_sim_m = z_fin_nom_m + ecart_nominal_mm/1000.0

    mean_z_fin_m = float(np.mean(z_fin_sim_m))
    std_z_fin_m = float(np.std(z_fin_sim_m, ddof=1)) if n_draws_val > 1 else 0.0
    std_z_fin_mm = std_z_fin_m * 1000.0

    quantiles_probs = [0.025, 0.05, 0.25, 0.50, 0.75, 0.95, 0.975]
    z_quantiles_m = np.quantile(z_fin_sim_m, quantiles_probs)
    ecart_quantiles_mm = np.quantile(ecart_nominal_mm, quantiles_probs)

    quantiles_dict_m = {
        f"P{str(round(p * 100, 1)).rstrip('0').rstrip('.')}": float(q)
        for p, q in zip(quantiles_probs, z_quantiles_m)
    }
    quantiles_ecart_mm = {
        f"P{str(round(p * 100, 1)).rstrip('0').rstrip('.')}": float(q)
        for p, q in zip(quantiles_probs, ecart_quantiles_mm)
    }

    has_reference = base["can_compensate"] and base["z_ref_used_m"] is not None
    draws_Index = np.arange(1, n_draws_val + 1, dtype=int)

    if has_reference:
        z_ref_used_m = float(base["z_ref_used_m"])
        f_nom_mm = float(base["fermeture_mm"])
        tol_mm = float(tolerance_mm)

        f_sim_mm = f_nom_mm + ecart_nominal_mm
        f_sim_m = f_sim_mm/1000
        exceed_mask = exceeds_tolerance(f_sim_mm,tol_mm)
        k_exceed = int(np.sum(exceed_mask))
        prob_exceed = float(k_exceed / n_draws_val)
        ci_low, ci_high = binomial_ci_95(k_exceed, n_draws_val)

        if sigma_total_analytic_mm > 0.0:
            p_within = stats.norm.cdf(
                tol_mm, loc=f_nom_mm, scale=sigma_total_analytic_mm
            ) - stats.norm.cdf(
                -tol_mm, loc=f_nom_mm, scale=sigma_total_analytic_mm
            )
            prob_analytic = float(max(0.0, min(1.0, 1.0 - p_within)))
        else:
            prob_analytic = float(exceeds_tolerance(f_nom_mm,tol_mm))

        cum_exceed = np.cumsum(exceed_mask.astype(float))
        running_prob = cum_exceed / draws_Index
        running_ci_low, running_ci_high = wilson_ci_95_vectorized(cum_exceed, draws_Index.astype(float))
        f_quantiles_mm = {
            f"P{str(round(p * 100, 1)).rstrip('0').rstrip('.')}": float(q)
            for p, q in zip(quantiles_probs, np.quantile(f_sim_mm, quantiles_probs))
        }
    else:
        f_sim_m = None
        f_sim_mm = None
        k_exceed = None
        prob_exceed = None
        ci_low, ci_high = (None, None)
        prob_analytic = None
        running_prob = None
        running_ci_low, running_ci_high = (None, None)
        f_quantiles_mm = None

    draws_df = pd.DataFrame(
        {
            "tirage_id": draws_Index,
            "z_fin_sim_m": z_fin_sim_m,
            "ecart_nominal_mm": ecart_nominal_mm,
        }
    )
    if has_reference and f_sim_mm is not None:
        draws_df["fermeture_sim_m"] = f_sim_m
        draws_df["fermeture_sim_mm"] = f_sim_mm
        draws_df["depasse_tolerance"] = exceeds_tolerance(f_sim_mm,tolerance_mm)

    return {
        "is_valid": True,
        "base": base,
        "sigma_mm": sigma_mm_val,
        "sigma_m": sigma_m_val,
        "n_draws": n_draws_val,
        "seed": seed_val,
        "n_stations": n_stations,
        "n_lectures": n_lectures,
        "var_station_mm2": var_station_mm2,
        "sigma_station_mm": sigma_station_mm,
        "var_total_mm2": var_total_mm2,
        "sigma_total_analytic_mm": sigma_total_analytic_mm,
        "sigma_total_analytic_m": sigma_total_analytic_m,
        "mean_z_fin_m": mean_z_fin_m,
        "std_z_fin_m": std_z_fin_m,
        "std_z_fin_mm": std_z_fin_mm,
        "rel_error_std_pct": (
            abs(std_z_fin_mm - sigma_total_analytic_mm) / sigma_total_analytic_mm * 100.0
            if sigma_total_analytic_mm > 0
            else 0.0
        ),
        "quantiles_z_m": quantiles_dict_m,
        "quantiles_ecart_mm": quantiles_ecart_mm,
        "has_reference": has_reference,
        "z_fin_sim_m": z_fin_sim_m,
        "ecart_nominal_mm": ecart_nominal_mm,
        "f_sim_m": f_sim_m,
        "f_sim_mm": f_sim_mm,
        "f_quantiles_mm": f_quantiles_mm,
        "k_exceed": k_exceed,
        "prob_exceed": prob_exceed,
        "ci_95_low": ci_low,
        "ci_95_high": ci_high,
        "prob_analytic": prob_analytic,
        "draws_index": draws_Index,
        "running_prob": running_prob,
        "running_ci_low": running_ci_low,
        "running_ci_high": running_ci_high,
        "draws_df": draws_df,
    }


def compare_two_precisions(
    stations_input: Any,
    z0: float,
    sigma_initial_mm: float,
    sigma_improved_mm: float,
    n_draws: int,
    seed: int,
    regime: str = "encadre",
    z_ref_final: Optional[float] = None,
    tolerance_mm: float = 15.0,
) -> Dict[str, Any]:
    """
    Compare deux précisions de lecture avec exactement les mêmes lectures nominales et la même graine.
    """
    sim_init = simulate_uncertainties(
        stations_input=stations_input,
        z0=z0,
        sigma_mm=sigma_initial_mm,
        n_draws=n_draws,
        seed=seed,
        regime=regime,
        z_ref_final=z_ref_final,
        tolerance_mm=tolerance_mm,
    )
    sim_imp = simulate_uncertainties(
        stations_input=stations_input,
        z0=z0,
        sigma_mm=sigma_improved_mm,
        n_draws=n_draws,
        seed=seed,
        regime=regime,
        z_ref_final=z_ref_final,
        tolerance_mm=tolerance_mm,
    )
    return {
        "initial": sim_init,
        "improved": sim_imp,
    }


# ==============================================================================
# 5. SCHÉMA PÉDAGOGIQUE SVG DYNAMIQUE D'UNE VISÉE DE NIVELLEMENT
# ==============================================================================

def build_sighting_svg(
    station_id: str,
    point_ar: str,
    point_av: str,
    lecture_ar_m: float,
    lecture_av_m: float,
    z_ar_m: float,
    longueur_m: Optional[float] = None,
) -> str:
    """
    Construit un schéma SVG pédagogique et proportionné illustrant la visée
    horizontale d'une station de nivellement géométrique.
    """
    full_title=escape(f'{station_id} : {point_ar} vers {point_av}')
    station_id,point_ar,point_av=(escape(str(v) if len(str(v))<=14 else str(v)[:11]+'...') for v in (station_id,point_ar,point_av))
    dz_m = float(lecture_ar_m) - float(lecture_av_m)
    z_av_m = float(z_ar_m) + dz_m
    hv_m = float(z_ar_m) + float(lecture_ar_m)

    y_sight = 105.0
    scale_px_per_m = min(48.0,155.0/max(float(lecture_ar_m),float(lecture_av_m),.001))

    l_ar_clamped = float(lecture_ar_m)
    l_av_clamped = float(lecture_av_m)

    y_ground_ar = y_sight + l_ar_clamped * scale_px_per_m
    y_ground_av = y_sight + l_av_clamped * scale_px_per_m
    y_ground_mid = 0.5 * (y_ground_ar + y_ground_av)

    staff_height_px = max(l_ar_clamped, l_av_clamped, 2.4) * scale_px_per_m + 22.0
    y_top_ar = max(28.0, y_ground_ar - staff_height_px)
    y_top_av = max(28.0, y_ground_av - staff_height_px)

    sens_label = (
        f"Montée (+{dz_m:.3f} m)"
        if dz_m > 1e-6
        else (f"Descente ({dz_m:.3f} m)" if dz_m < -1e-6 else "Terrain horizontal (0,000 m)")
    )
    dist_info = (
        f"Longueur du cheminement : {longueur_m:.1f} m (distincte des lectures sur mire)"
        if longueur_m is not None and not pd.isna(longueur_m) and longueur_m > 0
        else "Distance horizontale non renseignée (ΔZ dépend uniquement de L_AR et L_AV)"
    )

    x_ar = 135
    x_inst = 365
    x_av = 585
    x_cote = 745

    y_min_ground = min(y_ground_ar, y_ground_av)
    y_max_ground = max(y_ground_ar, y_ground_av)

    svg = f"""
    <div class="tc-svg-frame" style="background:#FFFFFF; border:1px solid #D5D9E0; border-radius:8px; padding:6px; margin-bottom:12px;overflow-x:auto;">
    <svg viewBox="0 0 840 356" width="840" height="356" xmlns="http://www.w3.org/2000/svg"
         style="width:100%; min-width:650px; height:auto; display:block; font-family:'Noto Sans','Segoe UI',sans-serif;">
      <title>{full_title}</title>
      <rect x="0" y="0" width="840" height="34" fill="#F3F4F6" rx="6"/>
      <text x="16" y="22" fill="#1E2229" font-size="13" font-weight="700">
        Schéma de visée — Station {station_id} ({point_ar} → {point_av}) • {sens_label}
      </text>
      <text x="824" y="22" fill="#4A515D" font-size="11.5" text-anchor="end">
        Plan de visée Hv = {hv_m:.3f} m
      </text>

      <!-- Sol / Profil du terrain naturel -->
      <path d="M 35 {y_ground_ar:.1f} L {x_ar} {y_ground_ar:.1f}
               C 245 {y_ground_ar:.1f}, 475 {y_ground_av:.1f}, {x_av} {y_ground_av:.1f}
               L 805 {y_ground_av:.1f} L 805 295 L 35 295 Z"
            fill="#EAECEF" stroke="none"/>
      <path d="M 35 {y_ground_ar:.1f} L {x_ar} {y_ground_ar:.1f}
               C 245 {y_ground_ar:.1f}, 475 {y_ground_av:.1f}, {x_av} {y_ground_av:.1f}
               L 805 {y_ground_av:.1f}"
            fill="none" stroke="#4A515D" stroke-width="2.5"/>

      <line x1="{x_ar}" y1="{y_ground_ar:.1f}" x2="{x_cote + 10}" y2="{y_ground_ar:.1f}"
            stroke="#94A3B8" stroke-width="1.3" stroke-dasharray="4,4"/>
      <line x1="{x_av}" y1="{y_ground_av:.1f}" x2="{x_cote + 10}" y2="{y_ground_av:.1f}"
            stroke="#94A3B8" stroke-width="1.3" stroke-dasharray="4,4"/>

      <!-- Mire Arrière (Framboise #B81D5A) -->
      <rect x="{x_ar - 7}" y="{y_top_ar:.1f}" width="14" height="{y_ground_ar - y_top_ar:.1f}"
            fill="#FFFFFF" stroke="#B81D5A" stroke-width="2"/>
      <line x1="{x_ar - 7}" y1="{y_sight:.1f}" x2="{x_ar + 7}" y2="{y_sight:.1f}"
            stroke="#B81D5A" stroke-width="3"/>
      <line x1="{x_ar - 24}" y1="{y_sight:.1f}" x2="{x_ar - 24}" y2="{y_ground_ar:.1f}"
            stroke="#B81D5A" stroke-width="2"/>
      <polygon points="{x_ar - 24},{y_sight:.1f} {x_ar - 28},{y_sight + 7:.1f} {x_ar - 20},{y_sight + 7:.1f}" fill="#B81D5A"/>
      <polygon points="{x_ar - 24},{y_ground_ar:.1f} {x_ar - 28},{y_ground_ar - 7:.1f} {x_ar - 20},{y_ground_ar - 7:.1f}" fill="#B81D5A"/>
      <rect x="16" y="{0.5 * (y_sight + y_ground_ar) - 18:.1f}" width="92" height="36" rx="4"
            fill="#FDF2F6" stroke="#B81D5A" stroke-width="1"/>
      <text x="62" y="{0.5 * (y_sight + y_ground_ar) - 3:.1f}" fill="#B81D5A" font-size="11" font-weight="700" text-anchor="middle">
        L_AR = {lecture_ar_m:.3f} m
      </text>
      <text x="62" y="{0.5 * (y_sight + y_ground_ar) + 11:.1f}" fill="#4A515D" font-size="10" text-anchor="middle">
        Lecture arrière
      </text>

      <!-- Mire Avant (Turquoise #0D9488) -->
      <rect x="{x_av - 7}" y="{y_top_av:.1f}" width="14" height="{y_ground_av - y_top_av:.1f}"
            fill="#FFFFFF" stroke="#0D9488" stroke-width="2"/>
      <line x1="{x_av - 7}" y1="{y_sight:.1f}" x2="{x_av + 7}" y2="{y_sight:.1f}"
            stroke="#0D9488" stroke-width="3"/>
      <line x1="{x_av - 24}" y1="{y_sight:.1f}" x2="{x_av - 24}" y2="{y_ground_av:.1f}"
            stroke="#0D9488" stroke-width="2"/>
      <polygon points="{x_av - 24},{y_sight:.1f} {x_av - 28},{y_sight + 7:.1f} {x_av - 20},{y_sight + 7:.1f}" fill="#0D9488"/>
      <polygon points="{x_av - 24},{y_ground_av:.1f} {x_av - 28},{y_ground_av - 7:.1f} {x_av - 20},{y_ground_av - 7:.1f}" fill="#0D9488"/>
      <rect x="{x_av - 122}" y="{0.5 * (y_sight + y_ground_av) - 18:.1f}" width="92" height="36" rx="4"
            fill="#F0FDFA" stroke="#0D9488" stroke-width="1"/>
      <text x="{x_av - 76}" y="{0.5 * (y_sight + y_ground_av) - 3:.1f}" fill="#0D9488" font-size="11" font-weight="700" text-anchor="middle">
        L_AV = {lecture_av_m:.3f} m
      </text>
      <text x="{x_av - 76}" y="{0.5 * (y_sight + y_ground_av) + 11:.1f}" fill="#4A515D" font-size="10" text-anchor="middle">
        Lecture avant
      </text>

      <!-- Niveau optique sur trépied au centre -->
      <line x1="{x_inst}" y1="{y_sight + 6:.1f}" x2="{x_inst - 22}" y2="{y_ground_mid:.1f}" stroke="#1E2229" stroke-width="2.5"/>
      <line x1="{x_inst}" y1="{y_sight + 6:.1f}" x2="{x_inst + 22}" y2="{y_ground_mid:.1f}" stroke="#1E2229" stroke-width="2.5"/>
      <line x1="{x_inst}" y1="{y_sight + 6:.1f}" x2="{x_inst}" y2="{y_ground_mid:.1f}" stroke="#1E2229" stroke-width="2"/>
      <rect x="{x_inst - 22}" y="{y_sight - 7:.1f}" width="44" height="14" rx="3" fill="#1E2229" stroke="#E59A00" stroke-width="1.5"/>

      <!-- Ligne de visée horizontale (Jaune signalétique #E59A00) -->
      <line x1="{x_ar}" y1="{y_sight:.1f}" x2="{x_av}" y2="{y_sight:.1f}"
            stroke="#E59A00" stroke-width="2.4" stroke-dasharray="7,4"/>
      <circle cx="{x_ar}" cy="{y_sight:.1f}" r="4.5" fill="#B81D5A"/>
      <circle cx="{x_av}" cy="{y_sight:.1f}" r="4.5" fill="#0D9488"/>
      <rect x="{x_inst - 95}" y="{y_sight - 30:.1f}" width="190" height="20" rx="4" fill="#FFFBEB" stroke="#E59A00" stroke-width="1"/>
      <text x="{x_inst}" y="{y_sight - 16:.1f}" fill="#1E2229" font-size="11" font-weight="700" text-anchor="middle">
        Ligne de visée horizontale (Niveau)
      </text>

      <!-- Cote de la dénivelée dZ à droite (bien dégagée du point avant) -->
      <line x1="{x_cote}" y1="{y_min_ground:.1f}" x2="{x_cote}" y2="{y_max_ground:.1f}"
            stroke="#1E2229" stroke-width="2"/>
      <rect x="{x_cote - 55}" y="{0.5 * (y_min_ground + y_max_ground) - 19:.1f}" width="126" height="38" rx="4"
            fill="#FFFFFF" stroke="#1E2229" stroke-width="1.3"/>
      <text x="{x_cote + 8}" y="{0.5 * (y_min_ground + y_max_ground) - 3:.1f}" fill="#1E2229" font-size="11" font-weight="700" text-anchor="middle">
        ΔZ = {dz_m:+.3f} m
      </text>
      <text x="{x_cote + 8}" y="{0.5 * (y_min_ground + y_max_ground) + 12:.1f}" fill="#4A515D" font-size="10" text-anchor="middle">
        L_AR − L_AV ({dz_m * 1000:+.1f} mm)
      </text>

      <!-- Repères au sol et altitudes -->
      <circle cx="{x_ar}" cy="{y_ground_ar:.1f}" r="5" fill="#1E2229"/>
      <text x="{x_ar}" y="{y_ground_ar + 18:.1f}" fill="#1E2229" font-size="11.2" font-weight="700" text-anchor="middle">
        {point_ar} (Z = {z_ar_m:.3f} m)
      </text>

      <circle cx="{x_av}" cy="{y_ground_av:.1f}" r="5" fill="#1E2229"/>
      <text x="{x_av}" y="{y_ground_av + 18:.1f}" fill="#1E2229" font-size="11.2" font-weight="700" text-anchor="middle">
        {point_av} (Z = {z_av_m:.3f} m)
      </text>

      <!-- Pied de schéma -->
      <rect x="0" y="296" width="840" height="60" fill="#F3F4F6"/>
      <text x="16" y="317" fill="#1E2229" font-size="11" font-weight="600">
        Z({point_av}) = Z({point_ar}) + (L_AR − L_AV) = {z_ar_m:.3f} + ({lecture_ar_m:.3f} − {lecture_av_m:.3f}) = {z_av_m:.3f} m
      </text>
      <text x="16" y="341" fill="#4A515D" font-size="10.5" text-anchor="start">
        {dist_info}
      </text>
    </svg>
    </div>
    """
    return svg.strip()


# ==============================================================================
# 6. GRAPHIQUES MATPLOTLIB AUX COULEURS COHÉRENTES
# ==============================================================================

def _apply_chart_style(ax: plt.Axes, title: str, xlabel: str, ylabel: str) -> None:
    """Applique le style visuel 'Carnet de terrain contemporain' à un axe Matplotlib."""
    ax.set_facecolor("#FFFFFF")
    ax.set_title(title, fontsize=11.5, fontweight="bold", color=PALETTE["anthracite"], pad=10, loc="left")
    ax.set_xlabel(xlabel, fontsize=9.8, fontweight="bold", color=PALETTE["anthracite"], labelpad=6)
    ax.set_ylabel(ylabel, fontsize=9.8, fontweight="bold", color=PALETTE["anthracite"], labelpad=6)
    ax.grid(True, linestyle="--", linewidth=0.7, color="#DCE0E6", alpha=0.85)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#94A3B8")
    ax.spines["bottom"].set_color("#94A3B8")
    ax.tick_params(colors=PALETTE["anthracite"], labelsize=9)


def plot_altimetric_profile(
    calc_result: Dict[str, Any],
    use_cumulative_distance: bool = True,
    figsize: Tuple[float, float] = (9.6, 4.3),
) -> plt.Figure:
    """
    Trace le profil altimétrique (élément visuel principal) :
    - Altitudes brutes en framboise (#B81D5A)
    - Altitudes compensées en turquoise (#0D9488) si applicables
    - Altitude finale de référence en jaune signalétique (#E59A00) si connue
    """
    fig, ax = plt.subplots(figsize=figsize, dpi=130)
    fig.patch.set_facecolor("#FFFFFF")

    pts = calc_result["points_table"]
    has_dist = bool(calc_result["all_lengths_valid"] and use_cumulative_distance)

    if has_dist:
        x_vals = pts["distance_cumulee_m"].to_numpy(dtype=float)
        xlabel = "Distance horizontale cumulée le long du cheminement (m)"
    else:
        x_vals = pts["rang"].to_numpy(dtype=float)
        xlabel = "Ordre séquentiel des points du cheminement (indice sans échelle de distance)"

    z_brut = pts["z_brut_m"].to_numpy(dtype=float)
    point_names = pts["point"].astype(str).tolist()

    all_z=z_brut
    if calc_result['can_compensate']:all_z=np.concatenate([all_z,pts['z_comp_m'].to_numpy(dtype=float)])
    if calc_result['z_ref_used_m'] is not None:all_z=np.append(all_z,calc_result['z_ref_used_m'])
    y_low,y_high=float(np.min(all_z)),float(np.max(all_z))
    y_floor = y_low - max(0.15, 0.20 * (y_high-y_low+0.1))
    ax.fill_between(x_vals, y_floor, z_brut, color=PALETTE["raspberry"], alpha=0.08, zorder=1)

    ax.plot(
        x_vals,
        z_brut,
        marker="o",
        markersize=7,
        linewidth=2.4,
        color=PALETTE["raspberry"],
        label="Profil altimétrique brut (mesures directes)",
        zorder=4,
    )

    if calc_result["can_compensate"] and pts["z_comp_m"].notna().all():
        z_comp = pts["z_comp_m"].to_numpy(dtype=float)
        ax.plot(
            x_vals,
            z_comp,
            marker="s",
            markersize=5.8,
            linewidth=2.0,
            linestyle="--",
            color=PALETTE["turquoise"],
            label="Profil altimétrique compensé",
            zorder=5,
        )

    if calc_result["z_ref_used_m"] is not None:
        z_ref = float(calc_result["z_ref_used_m"])
        ax.axhline(
            z_ref,
            color=PALETTE["signal_yellow"],
            linestyle=":",
            linewidth=1.9,
            label=f"Altitude finale de référence ({z_ref:.3f} m)",
            zorder=3,
        )
        ax.scatter(
            [x_vals[-1]],
            [z_ref],
            color=PALETTE["signal_yellow"],
            edgecolor=PALETTE["anthracite"],
            s=85,
            marker="D",
            zorder=6,
        )

    y_span = max(0.05,y_high-y_low)
    ax.set_ylim(y_floor, y_high + 0.52 * y_span + 0.08)

    # Detailed identifiers remain in the table; avoid label collisions on dense profiles.
    label_indices=set(range(len(pts))) if len(pts)<=5 else {0,len(pts)-1}
    for i, (x_i, z_i, name_i) in enumerate(zip(x_vals, z_brut, point_names)):
        if i not in label_indices:continue
        name_i=name_i if len(name_i)<=14 else name_i[:11]+'...'
        ax.annotate(
            f"{name_i}\n{z_i:.3f} m",
            xy=(x_i, z_i),
            xytext=(0, 9 if i % 2 == 0 else 12),
            textcoords="offset points",
            ha="center",
            va="bottom",
            fontsize=8.2,
            fontweight="bold",
            color=PALETTE["anthracite"],
            bbox=dict(boxstyle="round,pad=0.20", facecolor="#FFFFFF", edgecolor="#D5D9E0", alpha=0.92),
            zorder=7,
        )

    if not has_dist:
        ticks=np.unique(np.linspace(0,len(x_vals)-1,min(7,len(x_vals)),dtype=int))
        ax.set_xticks(x_vals[ticks])
        ax.set_xticklabels([f"Pt {int(x_vals[i])}" for i in ticks], rotation=0)

    _apply_chart_style(
        ax,
        title="Profil altimétrique du cheminement (brut vs compensé)",
        xlabel=xlabel,
        ylabel="Altitude Z (m)",
    )
    ax.legend(loc="upper left", frameon=True, facecolor="#FFFFFF", edgecolor="#D5D9E0", fontsize=8.5)
    fig.tight_layout()
    return fig


def plot_station_denivelees(
    calc_result: Dict[str, Any],
    figsize: Tuple[float, float] = (6.2, 3.6),
) -> plt.Figure:
    """Graphique en barres des dénivelées par station (distinguant montées et descentes)."""
    fig, ax = plt.subplots(figsize=figsize, dpi=130)
    fig.patch.set_facecolor("#FFFFFF")

    df = calc_result["stations_table"]
    labels = [f"{r['station_id']}\n({r['point_ar']}→{r['point_av']})" for _, r in df.iterrows()]
    dz_vals = df["denivelee_m"].to_numpy(dtype=float)
    bar_colors = [
        PALETTE["turquoise"] if v >= 0 else PALETTE["raspberry"]
        for v in dz_vals
    ]

    x = np.arange(len(labels))
    bars = ax.bar(x, dz_vals, width=0.52, color=bar_colors, edgecolor=PALETTE["anthracite"], linewidth=0.8, zorder=3)
    ax.axhline(0.0, color=PALETTE["anthracite"], linewidth=1.2, zorder=2)

    span = max(0.1, float(np.max(np.abs(dz_vals))))
    ax.set_ylim(min(0.0, float(np.min(dz_vals))) - 0.28 * span, max(0.0, float(np.max(dz_vals))) + 0.28 * span)

    for bar, val in zip(bars, dz_vals):
        y_pos = val + (0.04 * span if val >= 0 else -0.06 * span)
        va = "bottom" if val >= 0 else "top"
        ax.text(
            bar.get_x() + bar.get_width() / 2.0,
            y_pos,
            f"{val:+.3f} m\n({val * 1000:+.0f} mm)",
            ha="center",
            va=va,
            fontsize=8.2,
            fontweight="bold",
            color=PALETTE["anthracite"],
        )

    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=8.5)
    _apply_chart_style(
        ax,
        title="Dénivelées mesurées par station (ΔZ = L_AR − L_AV)",
        xlabel="Stations du cheminement",
        ylabel="Dénivelée ΔZ (m)",
    )
    fig.tight_layout()
    return fig


def plot_station_corrections(
    calc_result: Dict[str, Any],
    figsize: Tuple[float, float] = (6.2, 3.6),
) -> plt.Figure:
    """Graphique des corrections de compensation appliquées par station (en mm)."""
    fig, ax = plt.subplots(figsize=figsize, dpi=130)
    fig.patch.set_facecolor("#FFFFFF")

    df = calc_result["stations_table"]
    labels = [f"{r['station_id']}" for _, r in df.iterrows()]
    corr_mm = df["correction_mm"].to_numpy(dtype=float)

    x = np.arange(len(labels))
    bars = ax.bar(
        x,
        corr_mm,
        width=0.48,
        color=PALETTE["signal_yellow"],
        edgecolor=PALETTE["anthracite"],
        linewidth=0.9,
        zorder=3,
    )
    ax.axhline(0.0, color=PALETTE["anthracite"], linewidth=1.2, zorder=2)

    span = max(1.0, float(np.max(np.abs(corr_mm))))
    ax.set_ylim(min(0.0, float(np.min(corr_mm))) - 0.30 * span, max(0.0, float(np.max(corr_mm))) + 0.30 * span)

    for bar, val, (_, row) in zip(bars, corr_mm, df.iterrows()):
        y_pos = val + (0.05 * span if val >= 0 else -0.06 * span)
        va = "bottom" if val >= 0 else "top"
        ax.text(
            bar.get_x() + bar.get_width() / 2.0,
            y_pos,
            f"{val:+.2f} mm\n({row['part_poids_pct']:.1f} %)",
            ha="center",
            va=va,
            fontsize=8.2,
            fontweight="bold",
            color=PALETTE["anthracite"],
        )

    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    _apply_chart_style(
        ax,
        title="Corrections de compensation par station (en mm)",
        xlabel="Identifiant de station",
        ylabel="Correction c_i (mm)",
    )
    fig.tight_layout()
    return fig


def plot_uncertainty_histogram(
    sim_result: Dict[str, Any],
    figsize: Tuple[float, float] = (6.4, 3.9),
) -> plt.Figure:
    """
    Trace l'histogramme des fermetures simulées (si référence connue) ou des
    altitudes finales simulées (si cheminement ouvert), avec densité analytique.
    """
    fig, ax = plt.subplots(figsize=figsize, dpi=130)
    fig.patch.set_facecolor("#FFFFFF")

    has_ref = sim_result["has_reference"]
    sigma_analytic_mm = float(sim_result["sigma_total_analytic_mm"])

    if has_ref and sim_result["f_sim_mm"] is not None:
        data = sim_result["f_sim_mm"]
        f_nom_mm = float(sim_result["base"]["fermeture_mm"])
        tol_mm = float(sim_result["base"]["tolerance_mm"])

        if sigma_analytic_mm > 0:
            ax.hist(
                data,
                bins=36,
                density=True,
                color=PALETTE["raspberry"],
                alpha=0.28,
                edgecolor=PALETTE["raspberry"],
                label="Distribution Monte Carlo (fermeture f)",
                zorder=3,
            )
            x_grid = np.linspace(float(np.min(data)) - 2.0, float(np.max(data)) + 2.0, 250)
            y_pdf = stats.norm.pdf(x_grid, loc=f_nom_mm, scale=sigma_analytic_mm)
            ax.plot(
                x_grid,
                y_pdf,
                color=PALETTE["anthracite"],
                linewidth=2.0,
                label=f"Loi normale théorique N({f_nom_mm:+.1f} mm, {sigma_analytic_mm:.2f}² mm²)",
                zorder=4,
            )
        else:
            ax.axvline(
                f_nom_mm,
                color=PALETTE["raspberry"],
                linewidth=3.0,
                label=f"Valeur déterministe (σ=0) : f = {f_nom_mm:+.2f} mm",
            )

        ax.axvline(f_nom_mm, color=PALETTE["raspberry"], linestyle="--", linewidth=1.8, label=f"Fermeture nominale ({f_nom_mm:+.1f} mm)")
        ax.axvline(-tol_mm, color=PALETTE["turquoise"], linestyle=":", linewidth=2.0, label=f"Seuils de tolérance (±{tol_mm:.1f} mm)")
        ax.axvline(tol_mm, color=PALETTE["turquoise"], linestyle=":", linewidth=2.0)

        _apply_chart_style(
            ax,
            title="Distribution des fermetures simulées f = Z_sim − Z_réf",
            xlabel="Fermeture simulée f (mm)",
            ylabel="Densité de probabilité",
        )
    else:
        data_m = sim_result["z_fin_sim_m"]
        z_nom_m = float(sim_result["base"]["z_fin_brut_m"])
        sigma_analytic_m = float(sim_result["sigma_total_analytic_m"])

        if sigma_analytic_m > 0:
            ax.hist(
                data_m,
                bins=36,
                density=True,
                color=PALETTE["turquoise"],
                alpha=0.30,
                edgecolor=PALETTE["turquoise"],
                label="Distribution Monte Carlo de Z_final",
                zorder=3,
            )
            x_grid = np.linspace(float(np.min(data_m)), float(np.max(data_m)), 250)
            y_pdf = stats.norm.pdf(x_grid, loc=z_nom_m, scale=sigma_analytic_m)
            ax.plot(
                x_grid,
                y_pdf,
                color=PALETTE["anthracite"],
                linewidth=2.0,
                label=f"Densité analytique (σ = {sigma_analytic_mm:.2f} mm)",
                zorder=4,
            )
        else:
            ax.axvline(z_nom_m, color=PALETTE["turquoise"], linewidth=3.0, label=f"Z_final nominal = {z_nom_m:.3f} m")

        ax.axvline(z_nom_m, color=PALETTE["raspberry"], linestyle="--", linewidth=1.8, label=f"Résultat nominal ({z_nom_m:.3f} m)")
        _apply_chart_style(
            ax,
            title="Dispersion autour de l'altitude finale nominale (cheminement ouvert)",
            xlabel="Altitude finale simulée Z_final (m)",
            ylabel="Densité de probabilité",
        )

    ax.legend(loc="upper right", frameon=True, facecolor="#FFFFFF", edgecolor="#D5D9E0", fontsize=8.2)
    fig.tight_layout()
    return fig


def plot_probability_convergence(
    sim_result: Dict[str, Any],
    figsize: Tuple[float, float] = (6.4, 3.9),
) -> plt.Figure:
    """
    Trace la convergence de la probabilité estimée P(|f| > tolérance) avec son
    intervalle binomial à 95 %, ou la convergence de l'écart-type empirique si ouvert.
    """
    fig, ax = plt.subplots(figsize=figsize, dpi=130)
    fig.patch.set_facecolor("#FFFFFF")

    idx = sim_result["draws_index"]
    if sim_result["has_reference"] and sim_result["running_prob"] is not None:
        run_p = sim_result["running_prob"] * 100.0
        ci_l = sim_result["running_ci_low"] * 100.0
        ci_h = sim_result["running_ci_high"] * 100.0
        p_theo = float(sim_result["prob_analytic"]) * 100.0

        ax.fill_between(
            idx,
            ci_l,
            ci_h,
            color=PALETTE["signal_yellow"],
            alpha=0.25,
            label="Intervalle de confiance binomial à 95 %",
            zorder=2,
        )
        ax.plot(
            idx,
            run_p,
            color=PALETTE["raspberry"],
            linewidth=2.0,
            label=f"Probabilité estimée Monte Carlo ({run_p[-1]:.2f} %)",
            zorder=4,
        )
        ax.axhline(
            p_theo,
            color=PALETTE["turquoise"],
            linestyle="--",
            linewidth=1.8,
            label=f"Valeur analytique exacte ({p_theo:.2f} %)",
            zorder=3,
        )
        ax.set_ylim(-2.0, max(15.0, min(102.0, float(np.max(ci_h)) + 8.0)))
        _apply_chart_style(
            ax,
            title="Stabilisation de la probabilité P(|f_sim| > tolérance)",
            xlabel="Nombre de tirages Monte Carlo cumulés",
            ylabel="Probabilité de dépassement (%)",
        )
    else:
        ecarts = sim_result["ecart_nominal_mm"]
        cum_mean = np.cumsum(ecarts) / idx
        cum_sq = np.cumsum(ecarts**2) / idx
        cum_std = np.sqrt(np.maximum(0.0, cum_sq - cum_mean**2))
        theo_std = float(sim_result["sigma_total_analytic_mm"])

        ax.plot(
            idx,
            cum_std,
            color=PALETTE["turquoise"],
            linewidth=2.0,
            label=f"Écart-type empirique cumulé ({cum_std[-1]:.2f} mm)",
            zorder=4,
        )
        ax.axhline(
            theo_std,
            color=PALETTE["raspberry"],
            linestyle="--",
            linewidth=1.8,
            label=f"Écart-type analytique σ_Z = {theo_std:.2f} mm",
            zorder=3,
        )
        _apply_chart_style(
            ax,
            title="Convergence de l'écart-type empirique (cheminement ouvert)",
            xlabel="Nombre de tirages Monte Carlo cumulés",
            ylabel="Écart-type de l'altitude finale (mm)",
        )

    ax.legend(loc="upper right", frameon=True, facecolor="#FFFFFF", edgecolor="#D5D9E0", fontsize=8.2)
    fig.tight_layout()
    return fig


def plot_comparison_distributions(
    comp_result: Dict[str, Any],
    figsize: Tuple[float, float] = (9.2, 4.1),
) -> plt.Figure:
    """
    Superpose les distributions de la variante initiale (sigma_1) et de la
    variante améliorée (sigma_2) avec les mêmes mesures nominales et la même graine.
    """
    fig, ax = plt.subplots(figsize=figsize, dpi=130)
    fig.patch.set_facecolor("#FFFFFF")

    s_init = comp_result["initial"]
    s_imp = comp_result["improved"]
    has_ref = s_init["has_reference"]

    if has_ref and s_init["f_sim_mm"] is not None and s_imp["f_sim_mm"] is not None:
        d1 = s_init["f_sim_mm"]
        d2 = s_imp["f_sim_mm"]
        f_nom = float(s_init["base"]["fermeture_mm"])
        tol = float(s_init["base"]["tolerance_mm"])

        if s_init["sigma_mm"] > 0:
            ax.hist(
                d1,
                bins=38,
                density=True,
                color=PALETTE["raspberry"],
                alpha=0.26,
                edgecolor=PALETTE["raspberry"],
                label=f"Mesures initiales (σ_lecture = {s_init['sigma_mm']:.2f} mm → σ_Z = {s_init['sigma_total_analytic_mm']:.2f} mm)",
            )
        if s_imp["sigma_mm"] > 0:
            ax.hist(
                d2,
                bins=38,
                density=True,
                color=PALETTE["turquoise"],
                alpha=0.35,
                edgecolor=PALETTE["turquoise"],
                label=f"Variante améliorée (σ_lecture = {s_imp['sigma_mm']:.2f} mm → σ_Z = {s_imp['sigma_total_analytic_mm']:.2f} mm)",
            )

        ax.axvline(
            f_nom,
            color=PALETTE["anthracite"],
            linestyle="--",
            linewidth=2.2,
            label=f"Fermeture nominale commune ({f_nom:+.2f} mm)",
        )
        ax.axvline(
            0.0,
            color="#64748B",
            linestyle="-",
            linewidth=1.2,
            label="Fermeture nulle idéale (0 mm)",
        )
        ax.axvline(-tol, color=PALETTE["signal_yellow"], linestyle=":", linewidth=2.0, label=f"Tolérance (±{tol:.1f} mm)")
        ax.axvline(tol, color=PALETTE["signal_yellow"], linestyle=":", linewidth=2.0)

        _apply_chart_style(
            ax,
            title="Comparaison des distributions de fermeture (mêmes lectures nominales, même graine)",
            xlabel="Fermeture simulée f = Z_sim − Z_réf (mm)",
            ylabel="Densité de probabilité",
        )
    else:
        d1 = s_init["ecart_nominal_mm"]
        d2 = s_imp["ecart_nominal_mm"]
        if s_init["sigma_mm"] > 0:
            ax.hist(
                d1,
                bins=38,
                density=True,
                color=PALETTE["raspberry"],
                alpha=0.26,
                edgecolor=PALETTE["raspberry"],
                label=f"Mesures initiales (σ_lecture = {s_init['sigma_mm']:.2f} mm → σ_Z = {s_init['sigma_total_analytic_mm']:.2f} mm)",
            )
        if s_imp["sigma_mm"] > 0:
            ax.hist(
                d2,
                bins=38,
                density=True,
                color=PALETTE["turquoise"],
                alpha=0.35,
                edgecolor=PALETTE["turquoise"],
                label=f"Variante améliorée (σ_lecture = {s_imp['sigma_mm']:.2f} mm → σ_Z = {s_imp['sigma_total_analytic_mm']:.2f} mm)",
            )
        ax.axvline(0.0, color=PALETTE["anthracite"], linestyle="--", linewidth=2.0, label="Résultat nominal (écart 0 mm)")
        _apply_chart_style(
            ax,
            title="Comparaison de la dispersion autour de l'altitude nominale (cheminement ouvert)",
            xlabel="Écart simulé par rapport à l'altitude finale nominale (mm)",
            ylabel="Densité de probabilité",
        )

    ax.legend(loc="upper right", frameon=True, facecolor="#FFFFFF", edgecolor="#D5D9E0", fontsize=8.4)
    fig.tight_layout()
    return fig


# ==============================================================================
# 7. IMPORTS ET EXPORTS (MODÈLES, CSV, EXCEL .XLSX, JSON, RAPPORT PDF)
# ==============================================================================

def generate_template_csv_bytes() -> bytes:
    """Génère un modèle CSV vierge pré-structuré pour la saisie d'un carnet."""
    template_df = pd.DataFrame(
        [
            {
                "station_id": "S1",
                "point_ar": "RN_Depart",
                "point_av": "P1",
                "lecture_ar_m": 1.500,
                "lecture_av_m": 1.200,
                "longueur_m": 50.0,
            },
            {
                "station_id": "S2",
                "point_ar": "P1",
                "point_av": "RN_Arrivee",
                "lecture_ar_m": 1.350,
                "lecture_av_m": 1.150,
                "longueur_m": 50.0,
            },
        ]
    )
    return template_df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")


def generate_template_excel_bytes() -> bytes:
    """Génère un modèle Excel (.xlsx) propre avec notice d'utilisation."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Carnet_Saisie"

    headers = ["station_id", "point_ar", "point_av", "lecture_ar_m", "lecture_av_m", "longueur_m"]
    ws.append(headers)
    ws.append(["S1", "RN_Depart", "P1", 1.500, 1.200, 50.0])
    ws.append(["S2", "P1", "RN_Arrivee", 1.350, 1.150, 50.0])

    header_fill = PatternFill("solid", fgColor="B81D5A")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    ws_help = wb.create_sheet("Notice_Saisie")
    ws_help.append(["Colonne", "Description", "Unité / Règle"])
    ws_help.append(["station_id", "Identifiant unique de la station (ex: S1, S2)", "Texte non vide, unique"])
    ws_help.append(["point_ar", "Repère ou point visé en arrière", "Doit égaler le point_av précédent"])
    ws_help.append(["point_av", "Repère ou point visé en avant", "Texte non vide, distinct de point_ar"])
    ws_help.append(["lecture_ar_m", "Lecture arrière sur mire graduée", "Mètres (m), valeur >= 0"])
    ws_help.append(["lecture_av_m", "Lecture avant sur mire graduée", "Mètres (m), valeur >= 0"])
    ws_help.append(["longueur_m", "Longueur horizontale de la station (optionnelle)", "Mètres (m), > 0 si pondération"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_carnet_csv_bytes(calc_result: Dict[str, Any]) -> bytes:
    """Exporte le carnet complet (brut et compensé) en CSV UTF-8."""
    df = calc_result["stations_table"].copy()
    for col in df:
        if pd.api.types.is_string_dtype(df[col]) or df[col].dtype==object:
            df[col]=df[col].map(lambda v:"'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v)
    return df.to_csv(index=False, float_format="%.6f", encoding="utf-8-sig").encode("utf-8-sig")


def export_carnet_excel_bytes(
    calc_result: Dict[str, Any],
    project_meta: Dict[str, Any],
    sim_result: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Exporte un classeur Excel (.xlsx) structuré en 3 feuilles : Synthèse, Carnet, Points."""
    wb = Workbook()

    ws_syn = wb.active
    ws_syn.title = "Synthese"
    title_font = Font(name="Calibri", size=13, bold=True, color="1E2229")
    hdr_fill = PatternFill("solid", fgColor="B81D5A")
    hdr_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    border_thin = Border(
        left=Side(style="thin", color="D5D9E0"),
        right=Side(style="thin", color="D5D9E0"),
        top=Side(style="thin", color="D5D9E0"),
        bottom=Side(style="thin", color="D5D9E0"),
    )

    ws_syn["A1"] = "TopoCheck — Synthèse du Carnet de Nivellement Géométrique"
    ws_syn["A1"].font = title_font
    ws_syn.append([])
    ws_syn.append(["Paramètre / Indicateur", "Valeur", "Unité / Détail"])

    for col_i in range(1, 4):
        c = ws_syn.cell(row=3, column=col_i)
        c.fill = hdr_fill
        c.font = hdr_font

    ferm_str = f"{calc_result['fermeture_mm']:+.2f}" if calc_result["fermeture_mm"] is not None else "Non calculable (ouvert)"
    rows_syn = [
        ("Projet", project_meta.get("project_name", "TopoCheck"), ""),
        ("Opérateur / Groupe", project_meta.get("operator", "ENSPY 3GC"), ""),
        ("Régime du cheminement", calc_result["regime"], ""),
        ("Altitude de départ Z0", round(calc_result["z0_m"], 4), "m"),
        (
            "Altitude finale de référence",
            round(calc_result["z_ref_used_m"], 4) if calc_result["z_ref_used_m"] is not None else "Aucune",
            "m",
        ),
        ("Nombre de stations", calc_result["n_stations"], "stations"),
        ("Somme des lectures arrière (Σ L_AR)", round(calc_result["sum_ar_m"], 4), "m"),
        ("Somme des lectures avant (Σ L_AV)", round(calc_result["sum_av_m"], 4), "m"),
        ("Dénivelée totale brute (Σ ΔZ)", round(calc_result["sum_denivelees_m"], 4), "m"),
        ("Altitude finale brute calculée", round(calc_result["z_fin_brut_m"], 4), "m"),
        ("Contrôle arithmétique de carnet", "VALIDE" if calc_result["arithmetic_check_ok"] else "ERREUR", ""),
        ("Fermeture altimétrique f", ferm_str, "mm"),
        ("Tolérance pédagogique choisie", round(calc_result["tolerance_mm"], 2), "mm"),
        (
            "Altitude finale compensée",
            round(calc_result["z_fin_comp_m"], 4) if calc_result["z_fin_comp_m"] is not None else "Non applicable",
            "m",
        ),
    ]
    if sim_result and sim_result.get("is_valid"):
        rows_syn.extend(
            [
                ("Écart-type par lecture (σ_L)", round(sim_result["sigma_mm"], 2), "mm"),
                ("Écart-type analytique final (σ_Z)", round(sim_result["sigma_total_analytic_mm"], 2), "mm"),
                ("Écart-type Monte Carlo final", round(sim_result["std_z_fin_mm"], 2), "mm"),
            ]
        )

    for r_item in rows_syn:
        ws_syn.append(list(r_item))

    ws_syn.column_dimensions["A"].width = 36
    ws_syn.column_dimensions["B"].width = 24
    ws_syn.column_dimensions["C"].width = 22

    ws_car = wb.create_sheet("Carnet_Brut_Compense")
    df_car = calc_result["stations_table"].copy()
    car_cols = list(df_car.columns)
    ws_car.append(car_cols)
    for col_i in range(1, len(car_cols) + 1):
        c = ws_car.cell(row=1, column=col_i)
        c.fill = hdr_fill
        c.font = hdr_font
        ws_car.column_dimensions[get_column_letter(col_i)].width = 16

    for _, row in df_car.iterrows():
        vals = [None if pd.isna(v) else (round(float(v), 5) if isinstance(v, (float, np.floating)) else v) for v in row]
        ws_car.append(vals)

    ws_pts = wb.create_sheet("Profil_Points")
    df_pts = calc_result["points_table"].copy()
    pts_cols = list(df_pts.columns)
    ws_pts.append(pts_cols)
    for col_i in range(1, len(pts_cols) + 1):
        c = ws_pts.cell(row=1, column=col_i)
        c.fill = hdr_fill
        c.font = hdr_font
        ws_pts.column_dimensions[get_column_letter(col_i)].width = 18

    for _, row in df_pts.iterrows():
        vals = [None if pd.isna(v) else (round(float(v), 5) if isinstance(v, (float, np.floating)) else v) for v in row]
        ws_pts.append(vals)

    for sheet in (ws_syn, ws_car, ws_pts):
        for r in sheet.iter_rows():
            for cell in r:
                cell.border = border_thin
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.data_type=='f':cell.data_type='s'
        sheet.freeze_panes='A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_project_json_bytes(
    project_meta: Dict[str, Any],
    calc_result: Dict[str, Any],
    sim_params: Dict[str, Any],
) -> bytes:
    """Exporte l'intégralité du projet en JSON."""
    df_clean = calc_result["stations_table"][ALL_COLUMNS].copy()
    records = []
    for _, r in df_clean.iterrows():
        records.append(
            {
                "station_id": str(r["station_id"]),
                "point_ar": str(r["point_ar"]),
                "point_av": str(r["point_av"]),
                "lecture_ar_m": float(r["lecture_ar_m"]),
                "lecture_av_m": float(r["lecture_av_m"]),
                "longueur_m": None if pd.isna(r["longueur_m"]) else float(r["longueur_m"]),
            }
        )

    payload = {
        "application": "TopoCheck",
        "version": "1.0.0",
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "project_name": project_meta.get("project_name", "Projet TopoCheck"),
        "operator": project_meta.get("operator", "Étudiant 3GC ENSPY"),
        "data_origin": project_meta.get("data_origin", "Origine non documentée"),
        "parameters": {
            "z0": calc_result["z0_m"],
            "regime": calc_result["regime"],
            "z_ref_final": sim_params.get('z_ref_final',calc_result["z_ref_used_m"]),
            "tolerance_mm": calc_result["tolerance_mm"],
            "comp_mode": calc_result["comp_mode_requested"],
            "sigma_mm": sim_params.get("sigma_mm", 2.0),
            "sigma_alt_mm": sim_params.get("sigma_alt_mm", 0.8),
            "n_draws": sim_params.get("n_draws", 3000),
            "seed": sim_params.get("seed", 42),
        },
        "stations": records,
        "summary_results": {
            "sum_ar_m": calc_result["sum_ar_m"],
            "sum_av_m": calc_result["sum_av_m"],
            "sum_denivelees_m": calc_result["sum_denivelees_m"],
            "z_fin_brut_m": calc_result["z_fin_brut_m"],
            "fermeture_m": calc_result["fermeture_m"],
            "fermeture_mm": calc_result["fermeture_mm"],
            "z_fin_comp_m": calc_result["z_fin_comp_m"],
        },
    }
    return json.dumps(payload, indent=2, ensure_ascii=False,allow_nan=False).encode("utf-8")


def parse_uploaded_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Lit et valide un fichier importé (CSV, Excel .xlsx ou JSON projet).
    """
    fname = filename.lower().strip()
    try:
        if len(file_bytes)>2*1024*1024:raise ValueError('Fichier limité à 2 Mo.')
        if fname.endswith(".json"):
            data = json.loads(file_bytes.decode("utf-8-sig"))
            if not isinstance(data,dict) or "stations" not in data or not isinstance(data["stations"], list):
                return {
                    "ok": False,
                    "error": "Le fichier JSON ne contient pas de liste 'stations' valide.",
                }
            if data.get('version')==1 and 'depart_m' in data:
                if data['sigma_ar_mm']!=data['sigma_av_mm']:raise ValueError('Ancien projet avec sigmas AR/AV distincts : cette version utilise un sigma commun ; adaptation explicite nécessaire.')
                data={'project_name':data['titre'],'operator':'','stations':[dict(row,station_id=row['station']) for row in data['stations']],
                      'parameters':dict(z0=data['depart_m'],regime={'Liaison':'encadre','Boucle':'boucle','Ouvert':'ouvert'}[data['mode']],
                          z_ref_final=data['reference_m'],tolerance_mm=data['tolerance_mm'],comp_mode={'Égaux':'egaux','Longueurs':'longueurs'}[data['ponderation']],
                          sigma_mm=data['sigma_ar_mm'],sigma_alt_mm=data['sigma_ar_mm']/2,n_draws=data['tirages'],seed=data['graine'])}
            if data.get('application','TopoCheck')!='TopoCheck':raise ValueError('Ce projet ne provient pas de TopoCheck.')
            if not isinstance(data.get('parameters'),dict):raise ValueError('Paramètres JSON manquants ou invalides.')
            candidate=validated_project(dict(project_name=data.get('project_name','Projet importé'),operator=data.get('operator',''),
                                            data_origin=data.get('data_origin','Projet importé'),params=data['parameters'],stations=data['stations']))
            df = normalize_stations_df(candidate['stations'])
            params = candidate['params']
            return {
                "ok": True,
                "source_type": "json",
                "df": df,
                "params": params,
                "project_name": data.get("project_name"),
                "operator": data.get("operator"),
                "data_origin": candidate['data_origin'],
            }

        elif fname.endswith(".csv"):
            df_raw = pd.read_csv(io.BytesIO(file_bytes),dtype={c:str for c in ('station_id','point_ar','point_av')},nrows=101)
        elif fname.endswith(".xlsx"):
            wb = load_workbook(io.BytesIO(file_bytes), data_only=True,read_only=True)
            sheet_name = (
                "Carnet_Saisie"
                if "Carnet_Saisie" in wb.sheetnames
                else ("Carnet_Brut_Compense" if "Carnet_Brut_Compense" in wb.sheetnames else wb.sheetnames[0])
            )
            wb.close()
            df_raw = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_name,dtype={c:str for c in ('station_id','point_ar','point_av')},nrows=101)
        else:
            return {
                "ok": False,
                "error": f"Format de fichier non pris en charge ('{filename}'). Utilisez .csv, .xlsx ou .json.",
            }

        missing_cols = [c for c in REQUIRED_COLUMNS if c not in df_raw.columns]
        if missing_cols:
            return {
                "ok": False,
                "error": (
                    "Colonnes obligatoires manquantes dans le fichier importé : "
                    + ", ".join(missing_cols)
                    + f". Colonnes attendues : {', '.join(REQUIRED_COLUMNS)}."
                ),
            }

        df = normalize_stations_df(df_raw)
        if df.attrs.get('input_errors'):raise ValueError(' '.join(df.attrs['input_errors']))
        if len(df)>100:raise ValueError('Le carnet est limité à 100 stations.')
        return {
            "ok": True,
            "source_type": "table",
            "df": df,
            "params": {},
        }
    except Exception as exc:
        return {
            "ok": False,
            "error": f"Impossible de lire le fichier '{filename}' : {exc}",
        }


def generate_pdf_report_bytes(
    project_meta: Dict[str, Any],
    calc_result: Dict[str, Any],
    sim_result: Dict[str, Any],
) -> bytes:
    """
    Génère un rapport PDF complet et professionnel avec ReportLab.
    """
    fonts=Path(matplotlib.get_data_path())/'fonts'/'ttf'
    for name,file in [('TCsans','DejaVuSans.ttf'),('TCsans-Bold','DejaVuSans-Bold.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,str(fonts/file)))
    pdfmetrics.registerFontFamily('TCsans',normal='TCsans',bold='TCsans-Bold',italic='TCsans',boldItalic='TCsans-Bold')
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title="Rapport Topographique TopoCheck",
    )

    styles = getSampleStyleSheet()
    style_title = ParagraphStyle(
        "TC_Title",
        parent=styles["Heading1"],
        fontName="TCsans-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#B81D5A"),
        spaceAfter=4,
    )
    style_sub = ParagraphStyle(
        "TC_Sub",
        parent=styles["Normal"],
        fontName="TCsans",
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#4A515D"),
        spaceAfter=10,
    )
    style_h2 = ParagraphStyle(
        "TC_H2",
        parent=styles["Heading2"],
        fontName="TCsans-Bold",
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor("#1E2229"),
        spaceBefore=10,
        spaceAfter=5,
        keepWithNext=True,
    )
    style_body = ParagraphStyle(
        "TC_Body",
        parent=styles["Normal"],
        fontName="TCsans",
        fontSize=9.2,
        leading=13,
        textColor=colors.HexColor("#1E2229"),
        spaceAfter=6,
    )

    story: List[Any] = []
    cell_style=ParagraphStyle('TC_Cell',parent=style_body,fontSize=7.5,leading=10,spaceAfter=0)
    cell_header=ParagraphStyle('TC_CellHeader',parent=cell_style,textColor=colors.white,fontName='TCsans-Bold')
    def cells(rows):
        return [[Paragraph(escape(str(v)),cell_header if index==0 else cell_style) for v in row] for index,row in enumerate(rows)]

    story.append(Paragraph("TopoCheck — Rapport de Nivellement Géométrique Direct", style_title))
    story.append(
        Paragraph(
            f"Projet : <b>{escape(str(project_meta.get('project_name', 'Levé Topographique')))}</b> | "
            f"Opérateur : <b>{escape(str(project_meta.get('operator', 'Non renseigné')))}</b> | "
            f"Date : {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            style_sub,
        )
    )
    story.append(Paragraph('Origine : '+escape(str(project_meta.get('data_origin','Origine non documentée'))),style_body))

    story.append(Paragraph("1. Paramètres du levé, unités et synthèse des contrôles", style_h2))
    story.append(
        Paragraph(
            "Toutes les lectures sur mire, dénivelées et altitudes sont exprimées en <b>mètres (m)</b>. "
            "Les fermetures, tolérances, corrections et écarts-types de lecture sont également indiqués en "
            "<b>millimètres (mm)</b> (1 m = 1000 mm).",
            style_body,
        )
    )

    regime_labels = {
        "encadre": "Cheminement encadré entre deux repères connus",
        "boucle": "Boucle fermée revenant au repère initial",
        "ouvert": "Cheminement ouvert sans altitude finale connue",
    }
    ferm_txt = (
        f"{calc_result['fermeture_m']:+.4f} m ({calc_result['fermeture_mm']:+.1f} mm)"
        if calc_result["fermeture_m"] is not None
        else "Non calculable (cheminement ouvert sans référence finale)"
    )
    verdict_txt = (
        "Conforme à la tolérance pédagogique choisie"
        if calc_result["is_within_tolerance"] is True
        else (
            "Dépasse la tolérance pédagogique choisie"
            if calc_result["is_within_tolerance"] is False
            else "Non évaluable (référence absente)"
        )
    )

    syn_data = [
        ["Indicateur", "Valeur", "Observation / Contrôle"],
        ["Régime du cheminement", regime_labels.get(calc_result["regime"], calc_result["regime"]), f"{calc_result['n_stations']} station(s)"],
        ["Altitude de départ Z0", f"{calc_result['z0_m']:.3f} m", "Repère initial"],
        [
            "Altitude finale de référence",
            f"{calc_result['z_ref_used_m']:.3f} m" if calc_result["z_ref_used_m"] is not None else "Non renseignée",
            "Cible d'arrivée",
        ],
        [
            "Somme AR − Somme AV",
            f"{calc_result['sum_ar_m']:.3f} − {calc_result['sum_av_m']:.3f} = {calc_result['diff_lectures_m']:+.3f} m",
            "Contrôle arithmétique OK" if calc_result["arithmetic_check_ok"] else "Erreur arithmétique",
        ],
        ["Altitude finale brute calculée", f"{calc_result['z_fin_brut_m']:.3f} m", f"ΔZ total = {calc_result['sum_denivelees_m']:+.3f} m"],
        ["Fermeture altimétrique f", ferm_txt, f"Tolérance choisie : ±{calc_result['tolerance_mm']:.1f} mm"],
        ["Statut par rapport à la tolérance", verdict_txt, "Tolérance pédagogique (non normative)"],
        [
            "Altitude finale compensée",
            f"{calc_result['z_fin_comp_m']:.3f} m" if calc_result["z_fin_comp_m"] is not None else "Non compensé",
            f"Poids : {calc_result['comp_mode_effective']}" if calc_result["can_compensate"] else "Aucune compensation",
        ],
    ]

    t_syn = Table(cells(syn_data), colWidths=[52 * mm, 62 * mm, 64 * mm],repeatRows=1)
    t_syn.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#B81D5A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D5D9E0")),
                ("PADDING", (0, 0), (-1, -1), 4.5),
                ("VALIGN", (0,0),(-1,-1),"TOP"),
            ]
        )
    )
    story.append(t_syn)
    story.append(Spacer(1, 6))

    story.append(Paragraph("2. Carnet de nivellement détaillé (valeurs brutes et compensées)", style_h2))
    df_st = calc_result["stations_table"]
    carnet_rows = [
        [
            "Station",
            "Pt AR → Pt AV",
            "L_AR (m)",
            "L_AV (m)",
            "ΔZ brut (m)",
            "Z brut (m)",
            "Corr. (mm)",
            "Z comp. (m)",
        ]
    ]
    for _, r in df_st.iterrows():
        corr_cell = f"{r['correction_mm']:+.2f}" if not pd.isna(r["correction_mm"]) else "—"
        z_comp_cell = f"{r['z_av_comp_m']:.3f}" if not pd.isna(r["z_av_comp_m"]) else "—"
        carnet_rows.append(
            [
                str(r["station_id"]),
                f"{r['point_ar']} → {r['point_av']}",
                f"{r['lecture_ar_m']:.3f}",
                f"{r['lecture_av_m']:.3f}",
                f"{r['denivelee_m']:+.3f}",
                f"{r['z_av_brut_m']:.3f}",
                corr_cell,
                z_comp_cell,
            ]
        )

    t_car = Table(
        cells(carnet_rows),
        colWidths=[16 * mm, 34 * mm, 20 * mm, 20 * mm, 22 * mm, 22 * mm, 21 * mm, 23 * mm],
        repeatRows=1,
    )
    t_car.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E2229")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.0),
                ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D5D9E0")),
                ("PADDING", (0, 0), (-1, -1), 4.0),
                ("VALIGN", (0,0),(-1,-1),"TOP"),
            ]
        )
    )
    story.append(t_car)
    story.append(Spacer(1, 6))

    story.append(Paragraph("3. Analyse des incertitudes de lecture (Propagation & Monte Carlo)", style_h2))
    if sim_result and sim_result.get("is_valid"):
        prob_str = (
            f"{sim_result['prob_exceed'] * 100:.2f} % (IC 95 % : [{sim_result['ci_95_low'] * 100:.2f} %, {sim_result['ci_95_high'] * 100:.2f} %])"
            if sim_result["has_reference"] and sim_result["prob_exceed"] is not None
            else "Non applicable (cheminement ouvert sans référence finale)"
        )
        unc_rows = [
            ["Paramètre probabiliste", "Valeur", "Formule / Référence"],
            ["Écart-type par lecture sur mire (σ_L)", f"{sim_result['sigma_mm']:.2f} mm ({sim_result['sigma_m']:.5f} m)", "Loi normale centrée N(0, σ_L²)"],
            ["Nombre de tirages & graine", f"N = {sim_result['n_draws']} tirages | graine = {sim_result['seed']}", "Lectures saisies = référence du scénario"],
            [
                "Écart-type analytique final (σ_Z)",
                f"{sim_result['sigma_total_analytic_mm']:.2f} mm",
                f"σ_L × √(2n) avec 2n = {sim_result['n_lectures']} lectures",
            ],
            ["Écart-type empirique Monte Carlo", f"{sim_result['std_z_fin_mm']:.2f} mm", f"Moyenne Z_fin = {sim_result['mean_z_fin_m']:.4f} m"],
            [
                "Intervalle inter-quantiles [P2.5 ; P97.5]",
                f"[{sim_result['quantiles_z_m']['P2.5']:.4f} m ; {sim_result['quantiles_z_m']['P97.5']:.4f} m]",
                "95 % des altitudes finales simulées",
            ],
            ["Probabilité P(|f_sim| > tolérance)", prob_str, "Calculée uniquement si référence connue"],
        ]
        t_unc = Table(cells(unc_rows), colWidths=[56 * mm, 64 * mm, 58 * mm],repeatRows=1)
        t_unc.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D9488")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8.3),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D5D9E0")),
                    ("PADDING", (0, 0), (-1, -1), 4.0),
                    ("VALIGN", (0,0),(-1,-1),"TOP"),
                ]
            )
        )
        story.append(t_unc)

    story.append(PageBreak())
    for title,func,arg in [('Profil altimétrique',plot_altimetric_profile,calc_result),
                           ('Distribution simulée',plot_uncertainty_histogram,sim_result)]:
        if arg is None:continue
        fig=func(arg)
        try:
            image_buffer=io.BytesIO();fig.savefig(image_buffer,format='png',dpi=150,bbox_inches='tight');image_buffer.seek(0)
            width,height=fig.get_size_inches()
            story.append(KeepTogether([Paragraph(title,style_h2),Image(image_buffer,width=175*mm,height=175*mm*height/width),Spacer(1,12)]))
        finally:plt.close(fig)
    story.append(Paragraph("4. Rappel méthodologique et limites d'interprétation", style_h2))
    story.append(
        Paragraph(
            "• <b>Compensation mathématique vs validation terrain</b> : répartir la fermeture (−f) sur les stations "
            "permet de fermer algébriquement le carnet sur l'altitude de référence, mais n'élimine jamais une faute "
            "de lecture (ex. inversion d'un décimètre sur mire) et ne valide pas un levé hors tolérance.<br/>"
            "• <b>Modèle probabiliste</b> : la simulation Monte Carlo modélise uniquement les erreurs aléatoires "
            "gaussiennes centrées et indépendantes. Les biais systématiques (défaut de collimation du niveau, mire "
            "mal étalonnée, affaissement de crapaud) et les fautes grossières ne sont pas représentés.<br/>"
            "• <b>Tolérance pédagogique</b> : le seuil saisi par l'utilisateur sert à l'apprentissage du contrôle "
            "de fermeture et ne constitue pas une norme réglementaire officielle.",
            style_body,
        )
    )

    def footer(canvas,document):
        canvas.setFont('TCsans',8);canvas.drawString(16*mm,9*mm,'TopoCheck | Document pédagogique')
        canvas.drawRightString(A4[0]-16*mm,9*mm,str(document.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return buf.getvalue()


# ==============================================================================
# 8. SUITE DE TESTS DE RÉFÉRENCE INTÉGRÉE (SECTION 8 DU CAHIER DES CHARGES)
# ==============================================================================

def run_reference_test_suite() -> List[Dict[str, Any]]:
    """
    Exécute de manière déterministe tous les tests de référence exigés à la section 8.
    """
    results: List[Dict[str, Any]] = []

    t1_stations = [
        {"station_id": "S1", "point_ar": "RN0", "point_av": "RN1", "lecture_ar_m": 1.50, "lecture_av_m": 1.20, "longueur_m": 50.0}
    ]
    res1 = compute_leveling(t1_stations, z0=100.0, regime="encadre", z_ref_final=100.28, tolerance_mm=25.0, comp_mode="egaux")
    dz1 = float(res1["stations_table"].loc[0, "denivelee_m"])
    zfin1 = float(res1["z_fin_brut_m"])
    f1_m = float(res1["fermeture_m"])
    f1_mm = float(res1["fermeture_mm"])
    c1_mm = float(res1["stations_table"].loc[0, "correction_mm"])
    zcomp1 = float(res1["z_fin_comp_m"])
    ok1 = (
        abs(dz1 - 0.30) < 1e-9
        and abs(zfin1 - 100.30) < 1e-9
        and abs(f1_m - 0.02) < 1e-9
        and abs(f1_mm - 20.0) < 1e-6
        and abs(c1_mm - (-20.0)) < 1e-6
        and abs(zcomp1 - 100.28) < 1e-9
    )
    results.append(
        {
            "test": "1. Test simple (1 station, réf. 100,28 m)",
            "expected": "ΔZ = +0,300 m | Z_fin = 100,300 m | f = +0,020 m (+20,0 mm) | corr = −20,0 mm | Z_comp = 100,280 m",
            "actual": f"ΔZ = {dz1:+.3f} m | Z_fin = {zfin1:.3f} m | f = {f1_m:+.3f} m ({f1_mm:+.1f} mm) | corr = {c1_mm:+.1f} mm | Z_comp = {zcomp1:.3f} m",
            "passed": ok1,
        }
    )

    t2_stations = [
        {"station_id": "S1", "point_ar": "RN0", "point_av": "P1", "lecture_ar_m": 1.50, "lecture_av_m": 1.20, "longueur_m": 40.0},
        {"station_id": "S2", "point_ar": "P1", "point_av": "RN0", "lecture_ar_m": 1.10, "lecture_av_m": 1.40, "longueur_m": 60.0},
    ]
    res2 = compute_leveling(t2_stations, z0=100.0, regime="encadre", z_ref_final=100.0, tolerance_mm=10.0, comp_mode="egaux")
    zfin2 = float(res2["z_fin_brut_m"])
    f2_mm = float(res2["fermeture_mm"])
    ok2 = abs(zfin2 - 100.0) < 1e-9 and abs(f2_mm) < 1e-9
    results.append(
        {
            "test": "2. Test à deux stations (retour à 100,000 m)",
            "expected": "Z_fin = 100,000 m | fermeture f = 0,0 mm",
            "actual": f"Z_fin = {zfin2:.3f} m | fermeture f = {f2_mm:+.1f} mm",
            "passed": ok2,
        }
    )

    sim3_a = simulate_uncertainties(t2_stations, z0=100.0, sigma_mm=2.0, n_draws=2000, seed=777, regime="encadre", z_ref_final=100.0)
    sim3_b = simulate_uncertainties(t2_stations, z0=100.0, sigma_mm=2.0, n_draws=2000, seed=777, regime="encadre", z_ref_final=100.0)
    sim3_zero = simulate_uncertainties(t2_stations, z0=100.0, sigma_mm=0.0, n_draws=200, seed=777, regime="encadre", z_ref_final=100.0)

    sigma_theo = float(sim3_a["sigma_total_analytic_mm"])
    same_seed_identical = bool(np.allclose(sim3_a["z_fin_sim_m"], sim3_b["z_fin_sim_m"], atol=1e-12))
    zero_sigma_nominal = bool(np.allclose(sim3_zero["z_fin_sim_m"], 100.0, atol=1e-12))
    ok3 = abs(sigma_theo - 4.0) < 1e-9 and same_seed_identical and zero_sigma_nominal
    results.append(
        {
            "test": "3. Test probabiliste (2 stations, σ_L = 2 mm, σ = 0 mm et graine)",
            "expected": "σ_Z théorique = 4,00 mm | σ=0 → tous tirages = nominal | même graine → tirages identiques",
            "actual": f"σ_Z théorique = {sigma_theo:.2f} mm | σ=0 nominal={zero_sigma_nominal} | graine reproductible={same_seed_identical}",
            "passed": ok3,
        }
    )

    t4_stations = [
        {"station_id": "S1", "point_ar": "Haut", "point_av": "Bas", "lecture_ar_m": 0.650, "lecture_av_m": 2.150, "longueur_m": 30.0}
    ]
    res4 = compute_leveling(t4_stations, z0=50.0, regime="ouvert")
    dz4 = float(res4["stations_table"].loc[0, "denivelee_m"])
    zfin4 = float(res4["z_fin_brut_m"])
    ok4 = res4["is_valid"] and abs(dz4 - (-1.500)) < 1e-9 and abs(zfin4 - 48.500) < 1e-9
    results.append(
        {
            "test": "4. Test de descente (ΔZ < 0 accepté sans erreur)",
            "expected": "Valide = True | ΔZ = −1,500 m | Z_fin = 48,500 m",
            "actual": f"Valide = {res4['is_valid']} | ΔZ = {dz4:+.3f} m | Z_fin = {zfin4:.3f} m",
            "passed": ok4,
        }
    )

    sim5 = simulate_uncertainties(t4_stations, z0=50.0, sigma_mm=2.0, n_draws=500, seed=11, regime="ouvert", z_ref_final=None)
    ok5 = (
        sim5["base"]["fermeture_m"] is None
        and sim5["base"]["can_compensate"] is False
        and sim5["has_reference"] is False
        and sim5["prob_exceed"] is None
    )
    results.append(
        {
            "test": "5. Test sans référence finale (cheminement ouvert)",
            "expected": "fermeture = None | compensation = False | prob_exceed = None (incertitude autour du nominal seul)",
            "actual": f"fermeture = {sim5['base']['fermeture_m']} | compensation = {sim5['base']['can_compensate']} | prob_exceed = {sim5['prob_exceed']}",
            "passed": ok5,
        }
    )

    t6_stations = [
        {"station_id": "S1", "point_ar": "RN0", "point_av": "P1", "lecture_ar_m": 1.500, "lecture_av_m": 1.200, "longueur_m": 40.0},
        {"station_id": "S2", "point_ar": "P1", "point_av": "RN1", "lecture_ar_m": 1.120, "lecture_av_m": 1.400, "longueur_m": 60.0},
    ]
    res6 = compute_leveling(t6_stations, z0=100.0, regime="encadre", z_ref_final=100.000, tolerance_mm=25.0, comp_mode="longueurs")
    c6_1 = float(res6["stations_table"].loc[0, "correction_mm"])
    c6_2 = float(res6["stations_table"].loc[1, "correction_mm"])
    z6_comp = float(res6["z_fin_comp_m"])
    ok6 = abs(c6_1 - (-8.0)) < 1e-6 and abs(c6_2 - (-12.0)) < 1e-6 and abs(z6_comp - 100.000) < 1e-9
    results.append(
        {
            "test": "6. Test de compensation pondérée aux longueurs (40 m / 60 m, f = +20 mm)",
            "expected": "c1 = −8,00 mm (40 %) | c2 = −12,00 mm (60 %) | Σc = −20,00 mm | Z_comp = 100,000 m",
            "actual": f"c1 = {c6_1:+.2f} mm | c2 = {c6_2:+.2f} mm | Σc = {c6_1 + c6_2:+.2f} mm | Z_comp = {z6_comp:.3f} m",
            "passed": ok6,
        }
    )

    return results


# ==============================================================================
# 9. STYLES CSS « CARNET DE TERRAIN CONTEMPORAIN » & COMPOSANTS HTML
# ==============================================================================

def get_custom_css() -> str:
    """
    Feuille de style CSS intégrée à app.py.
    """
    return """
    <style>
      .stApp {
        background-color: #F3F4F6 !important;
        color: #1E2229 !important;
      }
      section[data-testid="stSidebar"] {
        background-color: #E8EAEE !important;
        border-right: 1px solid #CFD4DC !important;
      }

      /* En-tête compact type carnet d'opérateur */
      .tc-notebook-header {
        background-color: #FFFFFF;
        border: 1px solid #D5D9E0;
        border-left: 6px solid #B81D5A;
        border-radius: 6px;
        padding: 10px 16px;
        margin-bottom: 10px;
        display: flex;
        flex-wrap: wrap;
        justify-content: space-between;
        align-items: center;
        gap: 10px;
      }
      .tc-brand-title {
        font-size: 1.40rem;
        font-weight: 800;
        color: #1E2229;
        letter-spacing: -0.02em;
        margin: 0;
      }
      .tc-brand-badge {
        background-color: #B81D5A;
        color: #FFFFFF !important;
        font-size: 0.72rem;
        font-weight: 700;
        padding: 3px 8px;
        border-radius: 4px;
        margin-left: 8px;
        vertical-align: middle;
      }
      .tc-brand-sub {
        font-size: 0.84rem;
        color: #4A515D;
        margin: 2px 0 0 0;
      }
      .tc-core-question {
        background-color: #FFFBEB;
        border: 1px solid #E59A00;
        color: #1E2229;
        font-size: 0.83rem;
        font-weight: 600;
        padding: 5px 11px;
        border-radius: 5px;
      }

      /* Barre de synthèse horizontale compacte sur une seule ligne (non-quad-card) */
      .tc-summary-strip {
        background-color: #1E2229;
        color: #FFFFFF;
        border-radius: 6px;
        padding: 8px 14px;
        margin-bottom: 12px;
        display: flex;
        flex-wrap: nowrap;
        align-items: center;
        justify-content: space-between;
        gap: 8px;
        border-bottom: 3px solid #B81D5A;
      }
      .tc-strip-item {
        display: flex;
        flex-direction: column;
        padding-right: 10px;
        border-right: 1px solid #383F4C;
      }
      .tc-strip-item:last-child {
        border-right: none;
        padding-right: 0;
      }
      .tc-strip-label {
        font-size: 0.68rem;
        text-transform: uppercase;
        letter-spacing: 0.03em;
        color: #CBD5E1;
        margin-bottom: 2px;
      }
      .tc-strip-val {
        font-size: 0.88rem;
        font-weight: 700;
        color: #FFFFFF;
      }
      .tc-pill-ok {
        background-color: #0D9488;
        color: #FFFFFF !important;
        padding: 2px 7px;
        border-radius: 4px;
        font-size: 0.78rem;
        font-weight: 700;
        display: inline-block;
      }
      .tc-pill-warn {
        background-color: #E59A00;
        color: #1E2229 !important;
        padding: 2px 7px;
        border-radius: 4px;
        font-size: 0.78rem;
        font-weight: 700;
        display: inline-block;
      }
      .tc-pill-err {
        background-color: #B81D5A;
        color: #FFFFFF !important;
        padding: 2px 7px;
        border-radius: 4px;
        font-size: 0.78rem;
        font-weight: 700;
        display: inline-block;
      }

      /* Cartes de section style carnet de terrain */
      .tc-panel {
        background-color: #FFFFFF;
        border: 1px solid #D5D9E0;
        border-radius: 6px;
        padding: 11px 14px;
        margin-bottom: 10px;
        font-size: 0.87rem;
        line-height: 1.45;
      }
      .tc-panel-raspberry {
        background-color: #FDF2F6;
        border: 1px solid #F5C2D7;
        border-left: 5px solid #B81D5A;
        border-radius: 6px;
        padding: 10px 14px;
        margin-bottom: 10px;
        color: #1E2229;
        font-size: 0.87rem;
        line-height: 1.45;
      }
      .tc-panel-turquoise {
        background-color: #F0FDFA;
        border: 1px solid #99F6E4;
        border-left: 5px solid #0D9488;
        border-radius: 6px;
        padding: 10px 14px;
        margin-bottom: 10px;
        color: #1E2229;
        font-size: 0.87rem;
        line-height: 1.45;
      }
      .tc-panel-yellow {
        background-color: #FFFBEB;
        border: 1px solid #FDE68A;
        border-left: 5px solid #E59A00;
        border-radius: 6px;
        padding: 10px 14px;
        margin-bottom: 10px;
        color: #1E2229;
        font-size: 0.87rem;
        line-height: 1.45;
      }

      /* Rubrique active dans la barre latérale */
      div[role="radiogroup"] > label {
        background-color: #FFFFFF;
        border: 1px solid #D5D9E0;
        border-radius: 5px;
        padding: 6px 10px !important;
        margin-bottom: 5px !important;
      }

      /* Boutons contrastés framboise bien contenus */
      .stButton > button, .stDownloadButton > button, .tc-btn {
        background-color: #B81D5A !important;
        color: #FFFFFF !important;
        border: 1px solid #961548 !important;
        border-radius: 5px !important;
        font-weight: 600 !important;
        padding: 7px 12px !important;
        line-height: 1.35 !important;
        display: inline-block !important;
        text-align: center !important;
      }
      .stButton > button *, .stDownloadButton > button * {
        color: #FFFFFF !important;
      }

      @media (max-width: 900px) {
        .tc-summary-strip {
          flex-wrap: wrap;
          flex-direction: column;
          align-items: flex-start;
        }
        .tc-strip-item {
          border-right: none;
          border-bottom: 1px solid #383F4C;
          padding-bottom: 5px;
          width: 100%;
        }
      }
    </style>
    """


def render_header_and_summary_strip(
    project_meta: Dict[str, Any],
    calc_result: Dict[str, Any],
    sim_result: Optional[Dict[str, Any]],
) -> str:
    """Construit le code HTML de l'en-tête TopoCheck et de la barre de synthèse compacte."""
    project_meta={key:escape(str(value)) for key,value in project_meta.items()}
    header_html = f"""
    <div class="tc-notebook-header">
      <div>
        <h1 class="tc-brand-title">TopoCheck</h1>
        <span class="tc-brand-badge">CARNET ALTIMÉTRIQUE</span>
        <div class="tc-brand-sub">
          Nivellement géométrique direct • Calcul d'altitudes, fermeture &amp; incertitudes • <b>{project_meta.get('project_name', '')}</b>
        </div>
      </div>
      <div class="tc-core-question">
        🎯 « Mes lectures de nivellement permettent-elles de retrouver la bonne altitude ? »
      </div>
    </div>
    """

    if not calc_result.get("is_valid"):
        strip_html = """
        <div class="tc-summary-strip">
          <div class="tc-strip-item">
            <span class="tc-strip-label">État du carnet</span>
            <span class="tc-pill-err">Corrections requises dans la saisie</span>
          </div>
        </div>
        """
        return header_html + strip_html

    regime_short = {
        "encadre": "Encadré (2 repères)",
        "boucle": "Boucle fermée (retour Z0)",
        "ouvert": "Ouvert (sans réf. finale)",
    }.get(calc_result["regime"], calc_result["regime"])

    if calc_result["fermeture_mm"] is not None:
        f_mm = calc_result["fermeture_mm"]
        f_m = calc_result["fermeture_m"]
        tol = calc_result["tolerance_mm"]
        if calc_result["is_within_tolerance"]:
            ferm_badge = f'<span class="tc-pill-ok">|f| = {abs(f_mm):.1f} mm ≤ {tol:.1f} mm</span>'
        else:
            ferm_badge = f'<span class="tc-pill-err">|f| = {abs(f_mm):.1f} mm > {tol:.1f} mm</span>'
    else:
        ferm_badge = '<span class="tc-pill-warn">Non calculable (ouvert)</span>'

    sigma_txt = (
        f"&plusmn;{sim_result['sigma_total_analytic_mm']:.2f} mm (&sigma;<sub>L</sub>={sim_result['sigma_mm']:.1f} mm)"
        if sim_result and sim_result.get("is_valid")
        else "—"
    )

    strip_html = f"""
    <div class="tc-summary-strip">
      <div class="tc-strip-item">
        <span class="tc-strip-label">Cheminement &amp; Continuité</span>
        <span class="tc-strip-val">{regime_short} • {calc_result['n_stations']} st. • OK</span>
      </div>
      <div class="tc-strip-item">
        <span class="tc-strip-label">Départ Z0 &rarr; Arrivée brute</span>
        <span class="tc-strip-val">{calc_result['z0_m']:.3f} m &rarr; {calc_result['z_fin_brut_m']:.3f} m</span>
      </div>
      <div class="tc-strip-item">
        <span class="tc-strip-label">Dénivelée totale (&Sigma;L_AR &minus; &Sigma;L_AV)</span>
        <span class="tc-strip-val">{calc_result['sum_denivelees_m']:+.3f} m ({calc_result['sum_denivelees_m']*1000:+.1f} mm)</span>
      </div>
      <div class="tc-strip-item">
        <span class="tc-strip-label">Fermeture f vs Tolérance</span>
        <span>{ferm_badge}</span>
      </div>
      <div class="tc-strip-item">
        <span class="tc-strip-label">Écart-type analytique &sigma;_Z</span>
        <span class="tc-strip-val">{sigma_txt}</span>
      </div>
    </div>
    """
    return header_html + strip_html


# ==============================================================================
# 10. INITIALISATION DE LA SESSION ET APPLICATION PRINCIPALE STREAMLIT
# ==============================================================================

# APPLICATION STREAMLIT CORRIGEE
SECTIONS = ['Tableau de bord','Carnet de nivellement','Fermeture et compensation',
            'Incertitudes','Comparaison','Données et rapport','Méthode et tests']
SECTION_ICONS=['show_chart','edit_note','balance','query_stats','compare_arrows','download','functions']
PARAMETERS={'z0':100.,'regime':'encadre','z_ref_final':100.,'tolerance_mm':15.,'comp_mode':'egaux',
            'sigma_mm':2.,'sigma_alt_mm':.8,'n_draws':4000,'seed':42}


def new_project():
    return dict(project_name='Mon carnet',operator='',data_origin='Données utilisateur',params=copy.deepcopy(PARAMETERS),stations=[])


def validated_project(project):
    if not isinstance(project,dict) or not isinstance(project.get('params'),dict):
        raise ValueError('Projet et paramètres invalides.')
    p=project['params']
    if set(p)!=set(PARAMETERS):raise ValueError('Le projet doit contenir tous les paramètres de nivellement et de simulation.')
    for key,lo,hi in [('z0',-500,10000),('tolerance_mm',.1,1000),('sigma_mm',0,50),('sigma_alt_mm',0,50),('n_draws',100,50000),('seed',0,2147483647)]:
        if not finite_number(p[key],lo,hi):raise ValueError(f'{key} : nombre requis entre {lo} et {hi}.')
    if int(p['n_draws'])!=p['n_draws'] or int(p['seed'])!=p['seed']:raise ValueError('Tirages et graine : entiers requis.')
    if p['z_ref_final'] is not None and not finite_number(p['z_ref_final'],-500,10000):raise ValueError('Référence finale invalide.')
    for key in ('project_name','operator','data_origin'):
        if not isinstance(project.get(key,''),str) or len(project.get(key,''))>120:raise ValueError(f'{key} : texte limité à 120 caractères.')
    if not project.get('project_name','').strip():raise ValueError('Renseigner le nom du projet.')
    result=compute_leveling(project.get('stations'),**{k:p[k] for k in ('z0','regime','z_ref_final','tolerance_mm','comp_mode')})
    if not result['is_valid']:raise ValueError(' '.join(result['validation']['errors']))
    cleaned=copy.deepcopy(project)
    cleaned['stations']=result['stations_table'][ALL_COLUMNS].astype(object).where(pd.notna(result['stations_table'][ALL_COLUMNS]),None).to_dict('records')
    return cleaned


def project_from_preset(key):
    preset=get_preset_datasets()[key]
    return validated_project(dict(project_name=preset['project_name'],operator=preset['operator'],
        data_origin='Exemple fictif — aucune mesure réelle ENSPY',
        params={k:preset[k] for k in PARAMETERS},stations=preset['stations']))


def project_payload(project):
    return json.dumps(project,ensure_ascii=False,sort_keys=True,allow_nan=False)


def calculate_project(payload):
    c=json.loads(payload);p=c['params']
    return compute_leveling(c['stations'],**{k:p[k] for k in ('z0','regime','z_ref_final','tolerance_mm','comp_mode')})


def simulate_project(payload):
    c=json.loads(payload);p=c['params']
    return simulate_uncertainties(c['stations'],**{k:p[k] for k in ('z0','regime','z_ref_final','tolerance_mm','sigma_mm','n_draws','seed')})


def report_project(payload):
    c=json.loads(payload)
    return generate_pdf_report_bytes(c,calculate_project(payload),simulate_project(payload))


if st is not None:
    calculate_project=st.cache_data(max_entries=8,ttl=1800,show_spinner=False)(calculate_project)
    simulate_project=st.cache_data(max_entries=8,ttl=1800,show_spinner=False)(simulate_project)
    report_project=st.cache_data(max_entries=4,ttl=1800,show_spinner=False)(report_project)
    # st.html sanitizes inline SVG; CCv2 renders this escaped, generated diagram.
    sighting_component=st.components.v2.component('topocheck_sighting',
        html='<div id="sighting"></div>',
        css=':host {display:block;max-width:100%;} #sighting {max-width:100%;overflow-x:auto;}',
        js='export default function({data,parentElement}) {parentElement.querySelector("#sighting").innerHTML=data.html;}')


def commit_project(candidate):
    candidate=validated_project(candidate)
    st.session_state.tc_project=candidate
    st.session_state.tc_revision=st.session_state.get('tc_revision',0)+1


def load_selected_preset():
    commit_project(project_from_preset(st.session_state.tc_preset_choice))


def reset_project():
    st.session_state.tc_project=new_project()
    st.session_state.tc_revision=st.session_state.get('tc_revision',0)+1
    st.session_state.tc_view='Carnet de nivellement'


def navigate(section):st.session_state.tc_view=section


def interactive_chart(fig,key,height=360):
    fig.update_layout(template='plotly_white',height=height,paper_bgcolor='rgba(0,0,0,0)',
        font=dict(color='#30323c',size=12),colorway=['#B81D5A','#0D9488','#E59A00'],
        margin=dict(l=20,r=20,t=25,b=70),legend=dict(orientation='h',y=-.25))
    fig.update_xaxes(automargin=True);fig.update_yaxes(automargin=True)
    st.plotly_chart(fig,key=key,width='stretch',config={'displayModeBar':False})


def interactive_profile(calc):
    pts=calc['points_table'];x=pts.distance_cumulee_m if calc['all_lengths_valid'] else pts.rang
    fig=go.Figure(go.Scatter(x=x,y=pts.z_brut_m,text=pts.point,mode='lines+markers',name='Brut',
        hovertemplate='%{text}<br>Altitude %{y:.5f} m<extra></extra>'))
    if calc['can_compensate']:fig.add_trace(go.Scatter(x=x,y=pts.z_comp_m,text=pts.point,mode='lines+markers',name='Compensé',line=dict(dash='dot'),hovertemplate='%{text}<br>Altitude %{y:.5f} m<extra></extra>'))
    fig.update_layout(xaxis_title='Longueur cumulée (m)' if calc['all_lengths_valid'] else 'Ordre des points',yaxis_title='Altitude (m)')
    return fig


def revised_css():
    return '''<style>
    [data-testid="stMainBlockContainer"] {max-width:1450px;padding-top:4.2rem}
    h1,h2,h3,p,.tc-brand-title,.tc-strip-label {letter-spacing:0!important}
    h2 {font-size:1.4rem!important} h3 {font-size:1.15rem!important}
    .tc-brand-title {padding:0;line-height:1.3}
    .tc-notebook-header {border-radius:0;border-top:4px solid #B81D5A}
    .tc-panel {border-radius:0;box-shadow:none;border:0;border-bottom:1px solid #d5d9e0;background:transparent}
    .tc-summary-strip {border-radius:0;display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:14px}
    .tc-strip-item {min-width:0;overflow-wrap:anywhere}
    .tc-summary-strip .tc-strip-item:nth-child(n+4) {grid-column:span 1}
    .tc-svg-frame {max-width:100%;overflow-x:auto}
    button {border-radius:5px!important} button p {white-space:normal}
    [data-testid="stCaptionContainer"], [data-testid="stCaptionContainer"] p {color:#625b66!important;opacity:1!important}
    [data-testid="stSidebar"] .stButton > button {width:100%;min-height:40px}
    .stButton > button[data-testid="stBaseButton-secondary"] {background:#fff!important;color:#30323c!important;border-color:#c6cbd3!important}
    .stButton > button[data-testid="stBaseButton-secondary"] * {color:#30323c!important}
    .stButton > button:disabled {opacity:.45!important}
    .tc-pill-ok {background:#08776e}
    [data-testid="stForm"] {background:#fff;border-top:3px solid #B81D5A;border-radius:6px}
    @media(max-width:640px) {
       [data-testid="stMainBlockContainer"] {padding:4.2rem 1rem 1.5rem}
       [data-testid="stColumn"] {width:100%;flex:1 1 100%;min-width:0}
       .tc-summary-strip {grid-template-columns:1fr}
       .tc-notebook-header {padding:12px}
       .tc-brand-badge {display:inline-block;margin:5px 0}
       .tc-core-question {max-width:100%;overflow-wrap:anywhere}
       input {font-size:16px!important}
    }
    </style>'''


def editor_view(c,revision):
    st.caption('AR et AV sont des lectures sur mire en mètres. La longueur du cheminement est une autre grandeur. Les cellules ne sont appliquées qu’après Enregistrer.')
    with st.form(f'carnet_form_{revision}'):
        name=st.text_input('Nom du projet',c['project_name'],max_chars=120)
        operator=st.text_input('Opérateur ou groupe (facultatif)',c['operator'],max_chars=120)
        df=normalize_stations_df(c['stations'] if c['stations'] else pd.DataFrame(columns=ALL_COLUMNS))
        for col in ALL_COLUMNS[:3]:df[col]=df[col].astype('str')
        edited=st.data_editor(df,num_rows='dynamic',hide_index=True,width='stretch',key=f'carnet_editor_{revision}',
            column_config={k:st.column_config.NumberColumn(label,min_value=0.,max_value=10.,format='%.4f') for k,label in [('lecture_ar_m','Lecture AR (m)'),('lecture_av_m','Lecture AV (m)')]} |
                          {'longueur_m':st.column_config.NumberColumn('Longueur cheminement (m)',min_value=.001,max_value=10000.)})
        if st.form_submit_button('Enregistrer le carnet',icon=':material/save:'):
            draft=copy.deepcopy(c);clean=normalize_stations_df(edited)
            result=compute_leveling(clean,**{k:c['params'][k] for k in ('z0','regime','z_ref_final','tolerance_mm','comp_mode')})
            if not result['is_valid']:raise ValueError(' '.join(result['validation']['errors']))
            origin='Exemple fictif (éventuellement modifié)' if c['data_origin'].startswith('Exemple fictif') else 'Saisie utilisateur'
            draft.update(project_name=name.strip(),operator=operator,stations=clean.astype(object).where(pd.notna(clean),None).to_dict('records'),data_origin=origin)
            commit_project(draft);st.rerun()
    if c['stations']:
        calc=calculate_project(project_payload(c));table=calc['stations_table']
        i=st.selectbox('Station à illustrer',range(len(table)),format_func=lambda j:str(table.iloc[j].station_id))
        row=table.iloc[i]
        sighting_component(data={'html':build_sighting_svg(str(row.station_id),str(row.point_ar),str(row.point_av),row.lecture_ar_m,row.lecture_av_m,row.z_ar_brut_m,row.longueur_m)},key='sighting_diagram',height='content',width='stretch')
        st.caption('Schéma de visée sans échelle horizontale. Les noms complets et les valeurs sont conservés dans le carnet ; le dessin peut être défilé sur petit écran.')
        st.dataframe(table,hide_index=True,width='stretch')


def parameters_form(c,revision):
    p=c['params']
    with st.sidebar.expander('Paramètres du levé',expanded=not bool(c['stations'])):
        with st.form(f'parameters_{revision}'):
            z0=st.number_input('Altitude de départ (m)',-500.,10000.,float(p['z0']),.01,format='%.4f')
            regime=st.selectbox('Type de cheminement',['encadre','boucle','ouvert'],index=['encadre','boucle','ouvert'].index(p['regime']),format_func=lambda x:{'encadre':'Encadré — référence connue','boucle':'Boucle — retour au départ','ouvert':'Ouvert — sans référence'}[x])
            ref=st.number_input('Référence finale (m) — Encadré uniquement',-500.,10000.,float(p['z_ref_final'] if p['z_ref_final'] is not None else p['z0']),.01,format='%.4f')
            tol=st.number_input('Tolérance pédagogique (mm)',.1,1000.,float(p['tolerance_mm']),1.)
            if st.form_submit_button('Appliquer les paramètres',icon=':material/check:'):
                draft=copy.deepcopy(c);draft['params'].update(z0=z0,regime=regime,z_ref_final=ref if regime=='encadre' else None,tolerance_mm=tol)
                if c['stations']:commit_project(draft)
                else:
                    st.session_state.tc_project=draft;st.session_state.tc_revision=revision+1
                st.rerun()


def uncertainties_view(c,calc,sim,revision):
    p=c['params']
    with st.form(f'uncertainties_{revision}'):
        sigma=st.number_input('Écart-type par lecture (mm)',0.,50.,float(p['sigma_mm']),.2)
        n=st.number_input('Nombre de tirages',100,50000,int(p['n_draws']),100)
        seed=st.number_input('Graine aléatoire',0,2147483647,int(p['seed']))
        if st.form_submit_button('Appliquer et simuler',icon=':material/play_arrow:'):
            draft=copy.deepcopy(c);draft['params'].update(sigma_mm=sigma,n_draws=int(n),seed=int(seed));commit_project(draft);st.rerun()
    st.caption('Lectures nominales + erreurs normales indépendantes. Repères exacts dans le modèle, sans biais ni faute grossière. Les tirages ne sont pas des observations de terrain.')
    cols=st.columns(3)
    for col,label,value in zip(cols,['Sigma théorique','Sigma simulé','Hors tolérance'],[f"{sim['sigma_total_analytic_mm']:.3f} mm",f"{sim['std_z_fin_mm']:.3f} mm",f"{sim['prob_exceed']:.2%}" if sim['has_reference'] else 'Sans référence']):
        with col:st.metric(label,value)
    if sim['has_reference']:
        st.caption(f"IC binomial exact à 95 % : [{sim['ci_95_low']:.2%} ; {sim['ci_95_high']:.2%}]. Probabilité normale analytique : {sim['prob_analytic']:.2%}.")
        values=sim['f_sim_mm'];label='Fermeture brute (mm)'
    else:
        st.info('Sans référence finale : aucune fermeture ni probabilité de fermeture. Seule la dispersion autour de l’altitude nominale est estimée.')
        values=sim['ecart_nominal_mm'];label='Écart au résultat nominal (mm)'
    fig=go.Figure(go.Histogram(x=values,nbinsx=40,name='Tirages'))
    if sim['has_reference']:
        for bound in [-p['tolerance_mm'],p['tolerance_mm']]:fig.add_vline(x=bound,line_color='#0D9488',line_dash='dash')
    fig.update_layout(xaxis_title=label,yaxis_title='Nombre de tirages');interactive_chart(fig,'uncertainty_hist')
    if sim['has_reference']:
        idx=np.unique(np.linspace(0,len(values)-1,min(200,len(values)),dtype=int))
        fig=go.Figure(go.Scatter(x=sim['draws_index'][idx],y=sim['running_prob'][idx],name='Estimation'))
        fig.add_trace(go.Scatter(x=sim['draws_index'][idx],y=sim['running_ci_high'][idx],line=dict(width=0),showlegend=False))
        fig.add_trace(go.Scatter(x=sim['draws_index'][idx],y=sim['running_ci_low'][idx],line=dict(width=0),fill='tonexty',fillcolor='rgba(13,148,136,.12)',name='Wilson 95 % ponctuel'))
        fig.update_layout(xaxis_title='Tirages cumulés',yaxis_title='Probabilité estimée');interactive_chart(fig,'uncertainty_convergence')
        st.caption('L’intervalle final est exact ; la bande de convergence est une approximation de Wilson, ponctuelle et non simultanée.')
    st.dataframe(pd.DataFrame({'Quantile':list(sim['quantiles_z_m']),'Altitude finale (m)':list(sim['quantiles_z_m'].values()),'Erreur (mm)':list(sim['quantiles_ecart_mm'].values())}),hide_index=True,width='stretch')
    st.caption('Les quantiles décrivent les résultats, pas l’intervalle de confiance de la probabilité. Un écart-type n’est pas une erreur maximale garantie.')


def comparison_view(c,sim,revision):
    with st.form(f'comparison_{revision}'):
        sigma=st.number_input('Écart-type de la variante (mm)',0.,50.,float(c['params']['sigma_alt_mm']),.2)
        if st.form_submit_button('Comparer les précisions',icon=':material/compare_arrows:'):
            draft=copy.deepcopy(c);draft['params']['sigma_alt_mm']=sigma;commit_project(draft);st.rerun()
    alt=copy.deepcopy(c);alt['params']['sigma_mm']=alt['params']['sigma_alt_mm'];other=simulate_project(project_payload(alt))
    table=pd.DataFrame({'Scénario':['Actuel','Variante'],'Sigma lecture (mm)':[sim['sigma_mm'],other['sigma_mm']],
        'Sigma final (mm)':[sim['sigma_total_analytic_mm'],other['sigma_total_analytic_mm']],
        'P hors tolérance (%)':[sim['prob_exceed']*100 if sim['has_reference'] else None,other['prob_exceed']*100 if other['has_reference'] else None]})
    st.dataframe(table,hide_index=True,width='stretch')
    fig=go.Figure()
    for label,s in [('Actuel',sim),('Variante',other)]:fig.add_trace(go.Histogram(x=s['ecart_nominal_mm'],nbinsx=40,name=label,opacity=.55))
    fig.update_layout(barmode='overlay',xaxis_title='Erreur aléatoire finale (mm)',yaxis_title='Tirages');interactive_chart(fig,'comparison_hist')
    st.warning('Réduire sigma réduit la dispersion, pas la fermeture nominale. Si celle-ci dépasse déjà la tolérance, le risque de rester hors tolérance peut augmenter. Les deux scénarios utilisent des tirages appariés.')


def data_view(c,calc,sim):
    st.download_button('Modèle CSV',generate_template_csv_bytes(),file_name='topocheck-modele.csv')
    st.download_button('Modèle Excel',generate_template_excel_bytes(),file_name='topocheck-modele.xlsx')
    upload=st.file_uploader('Importer CSV, Excel ou JSON',type=['csv','xlsx','json'],max_upload_size=2)
    if upload:
        parsed=parse_uploaded_file(upload.getvalue(),upload.name)
        if not parsed['ok']:st.error(parsed['error'])
        else:
            st.dataframe(parsed['df'],hide_index=True,width='stretch')
            confirm=st.checkbox('Autoriser le remplacement des données',key='tc_confirm_import')
            if st.button('Appliquer le fichier importé',disabled=not confirm,icon=':material/upload:'):
                draft=copy.deepcopy(c)
                draft['stations']=parsed['df'].astype(object).where(pd.notna(parsed['df']),None).to_dict('records')
                if parsed['source_type']=='json':
                    draft['params']=parsed['params'];draft['project_name']=parsed.get('project_name') or 'Projet importé';draft['operator']=parsed.get('operator') or ''
                draft['data_origin']=parsed.get('data_origin','Données importées — origine à documenter')
                commit_project(draft);st.session_state.tc_flash='Projet importé avec tous ses paramètres.';st.rerun()
    if not calc or not calc['is_valid']:return
    st.subheader('Exports du projet enregistré')
    st.download_button('Projet JSON complet',export_project_json_bytes(c,calc,c['params']),file_name='topocheck-projet.json',mime='application/json')
    st.download_button('Carnet CSV',export_carnet_csv_bytes(calc),file_name='topocheck-carnet.csv')
    st.download_button('Classeur Excel',export_carnet_excel_bytes(calc,c,sim),file_name='topocheck-carnet.xlsx')
    st.download_button('Tirages CSV',sim['draws_df'].to_csv(index=False).encode('utf-8-sig'),file_name='topocheck-tirages.csv')
    if st.button('Préparer le rapport PDF',icon=':material/picture_as_pdf:'):
        with st.spinner('Préparation du rapport…'):st.session_state.tc_report=(project_payload(c),report_project(project_payload(c)))
    report=st.session_state.get('tc_report')
    if report and report[0]==project_payload(c):st.download_button('Télécharger le PDF',report[1],file_name='topocheck-rapport.pdf',mime='application/pdf')
    st.caption('CSV UTF-8, séparateur virgule, point décimal. Les imports tabulaires conservent les paramètres actuels. Le JSON restaure tous les paramètres. Les données restent en mémoire de session : téléchargez le JSON pour les conserver.')


def main():
    st.set_page_config(page_title='TopoCheck · Version corrigée',page_icon=':material/straighten:',layout='wide')
    st.html(get_custom_css());st.html(revised_css())
    st.session_state.setdefault('tc_project',new_project())
    c=copy.deepcopy(st.session_state.tc_project);revision=st.session_state.get('tc_revision',0)
    view=st.session_state.get('tc_view',SECTIONS[0])
    try:
        with st.sidebar:
            st.subheader('TopoCheck')
            for name,icon in zip(SECTIONS,SECTION_ICONS):st.button(name,icon=f':material/{icon}:',key=f'nav_{icon}',on_click=navigate,args=(name,),type='primary' if name==view else 'secondary',width='stretch')
            st.divider()
            presets=get_preset_datasets()
            labels=dict(chantier_enspy='Campus : 4 stations',test_simple_1_station='Test 1 : une station',test_2_stations_boucle='Tests 2 et 3 : boucle',cheminement_ouvert_descente='Ouvert : descente')
            st.selectbox('Exemple fictif',list(presets),format_func=lambda key:labels[key],key='tc_preset_choice')
            replace=st.checkbox('Confirmer le remplacement du carnet',key='tc_replace')
            st.button('Charger le scénario',icon=':material/science:',disabled=bool(c['stations']) and not replace,on_click=load_selected_preset,width='stretch')
            st.button('Nouveau carnet',icon=':material/note_add:',disabled=bool(c['stations']) and not replace,on_click=reset_project,width='stretch')
        parameters_form(c,revision)
        calc=calculate_project(project_payload(c)) if c['stations'] else None
        sim=simulate_project(project_payload(c)) if calc and calc['is_valid'] else None
        if calc:
            st.html(render_header_and_summary_strip(c,calc,sim))
        else:st.html('<header class="tc-notebook-header"><h1 class="tc-brand-title">TopoCheck</h1><p>Nivellement géométrique — carnet vide</p></header>')
        st.subheader(view)
        st.caption(c['project_name']+' · '+c['data_origin'])
        if 'tc_flash' in st.session_state:st.success(st.session_state.pop('tc_flash'))
        if view=='Carnet de nivellement':editor_view(c,revision)
        elif view=='Données et rapport':data_view(c,calc,sim)
        elif view=='Méthode et tests':
            st.markdown('**Dénivelée = AR − AV.** Altitude suivante = altitude précédente + dénivelée. Fermeture = arrivée brute − référence. La compensation répartit −f, sans corriger les fautes.')
            st.markdown('**Sigma final = sigma lecture × √(2 × nombre de stations).** Modèle indépendant, repères exacts ; aucune validation normative.')
            if st.button('Exécuter les tests de référence',icon=':material/task_alt:'):st.session_state.tc_checks=run_reference_test_suite()
            if 'tc_checks' in st.session_state:
                tests=st.session_state.tc_checks;st.dataframe(pd.DataFrame(tests),hide_index=True,width='stretch')
                if all(r['passed'] for r in tests):st.success(f'{len(tests)} tests de référence réussis.')
        elif not calc:
            st.info('Aucune donnée enregistrée. Chargez volontairement un scénario fictif ou créez votre carnet.')
            st.button('Créer mon carnet',icon=':material/add:',type='primary',on_click=navigate,args=('Carnet de nivellement',))
        elif not calc['is_valid']:
            st.error(' '.join(calc['validation']['errors']))
        elif view=='Fermeture et compensation':
            with st.form(f'compensation_{revision}'):
                mode=st.selectbox('Poids de compensation',['egaux','longueurs'],index=['egaux','longueurs'].index(c['params']['comp_mode']))
                if st.form_submit_button('Appliquer la compensation',icon=':material/balance:'):
                    draft=copy.deepcopy(c);draft['params']['comp_mode']=mode;commit_project(draft);st.rerun()
            if not calc['can_compensate']:st.info('Sans référence finale, ni fermeture ni compensation ne sont calculables.')
            else:
                st.warning('La compensation est pédagogique ; elle ne supprime pas une faute et ne certifie pas le levé.')
                st.dataframe(calc['stations_table'][['station_id','denivelee_m','correction_mm','denivelee_comp_m','z_av_brut_m','z_av_comp_m']],hide_index=True,width='stretch')
                interactive_chart(interactive_profile(calc),'comp_profile')
        elif view=='Incertitudes':uncertainties_view(c,calc,sim,revision)
        elif view=='Comparaison':comparison_view(c,sim,revision)
        else:
            interactive_chart(interactive_profile(calc),'dashboard_profile',410)
            if calc['fermeture_mm'] is not None:
                (st.success if calc['is_within_tolerance'] else st.warning)(f"Fermeture brute {calc['fermeture_mm']:+.2f} mm pour une tolérance pédagogique de ±{calc['tolerance_mm']:.1f} mm. Ce contrôle ne certifie pas le levé.")
            table=calc['stations_table']
            interactive_chart(go.Figure(go.Bar(x=table.station_id,y=table.denivelee_m)).update_layout(yaxis_title='Dénivelée (m)'),'dashboard_dn',280)
            st.dataframe(calc['points_table'],hide_index=True,width='stretch')
        if calc:
            for warning in calc['validation']['warnings']:st.caption(warning)
    except (ValueError,TypeError,KeyError,UnicodeError) as exc:
        st.error(f'Données non appliquées : {exc}')


if __name__=='__main__':main()
